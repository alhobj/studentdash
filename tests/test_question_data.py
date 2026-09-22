"""End-to-end coverage using only generated fictional workbooks."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from create_example_workbook import create_example
from studentdash.analytics import build_dashboard
from studentdash.config import Config
from studentdash.excel import WorkbookError, read_workbook
from studentdash.render import generate_dashboards
from studentdash.teacher import create_app


class QuestionDataTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='studentdash-question-test-')
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.path = self.root / 'fictional.xlsx'
        self.config = Config(self.path, self.root / 'output')
        create_example(self.path)

    def edit(self, change):
        book = load_workbook(self.path)
        try:
            change(book)
            book.save(self.path)
        finally:
            book.close()

    def test_personal_evidence_and_denominators(self):
        data = read_workbook(self.path)
        first = build_dashboard(data, '1001')
        second = build_dashboard(data, '1002')
        self.assertEqual(first.practice, [])
        self.assertAlmostEqual(first.overall_percent, 600 / 7)
        self.assertEqual(first.strengths, ['Explain'])
        self.assertEqual(second.work_on, ['Explain'])
        atomic = next(c for c in second.topics if c.name == 'Atomic structure')
        self.assertEqual((atomic.score, atomic.maximum, len(atomic.questions)), (2, 4, 2))
        self.assertFalse(any('differs from Results' in w for w in data.warnings))
        self.assertEqual(len(first.tickets), 1)
        self.assertNotEqual(first.tickets, second.tickets)

    def test_filter_updates_totals_evidence_resources_and_tickets(self):
        data = read_workbook(self.path)
        view = build_dashboard(data, '1001', assessment_id='A2')
        self.assertEqual((view.total_score, view.total_possible), (3, 3))
        self.assertEqual(len(view.history), 1)
        self.assertEqual(len(view.questions), 2)
        self.assertEqual([r.topic for r in view.resources], ['Stoichiometry'])
        self.assertEqual(view.tickets, [])
        self.assertEqual(view.strengths, [])  # Only one graded Explain question.

    def test_assessment_boundaries_override_global_only_for_that_assessment(self):
        self.edit(lambda b: setattr(b['Results']['C5'], 'value', 8))
        view = build_dashboard(read_workbook(self.path), '1002')
        self.assertEqual(next(h.grade for h in view.history if h.id == 'A2'), 6)
        self.assertEqual(next(h.grade for h in view.history if h.id == 'A1'), 4)

    def test_every_non_graded_status_excludes_marks(self):
        for status in ['missing', 'pending', 'absent', 'exempt']:
            with self.subTest(status=status):
                self.edit(lambda b: setattr(b['QuestionResults']['E7'], 'value', status))
                view = build_dashboard(read_workbook(self.path), '1001', assessment_id='A2')
                self.assertEqual((view.topics[0].score, view.topics[0].maximum), (3, 3))
                self.assertEqual(view.questions[1].status, status)

    def test_empty_evidence_stays_empty_with_examples_enabled(self):
        view = build_dashboard(read_workbook(self.path), '1001', assessment_id='A3')
        self.assertEqual(view.topics, [])
        self.assertEqual(view.practice, [])
        self.assertIsNone(view.overall_percent)

    def test_empty_normalized_tables_do_not_invent_demonstration_results(self):
        def clear(book):
            for name in ['Memberships', 'QuestionResults', 'Results', 'ExitTickets']:
                book[name].delete_rows(2, book[name].max_row)
        self.edit(clear)
        view = build_dashboard(read_workbook(self.path), '1001')
        self.assertEqual(view.practice, [])
        self.assertEqual(view.history, [])
        self.assertEqual(view.topics, [])

    def test_duplicate_question_result_rejected(self):
        self.edit(lambda b: b['QuestionResults'].append(['A1', 'Q1', '1001', 1, 'graded']))
        with self.assertRaisesRegex(WorkbookError, 'duplicate assessment/question/student'):
            read_workbook(self.path)

    def test_question_assessment_mismatch_rejected(self):
        self.edit(lambda b: setattr(b['QuestionResults']['B2'], 'value', 'Q3'))
        with self.assertRaisesRegex(WorkbookError, 'does not belong'):
            read_workbook(self.path)

    def test_hl_question_rejected_for_sl_member(self):
        self.edit(lambda b: b['QuestionResults'].append(['A2', 'Q5', '1001', 1, 'graded']))
        with self.assertRaisesRegex(WorkbookError, 'membership/level'):
            read_workbook(self.path)

    def test_membership_required(self):
        self.edit(lambda b: b.remove(b['Memberships']))
        with self.assertRaisesRegex(WorkbookError, 'requires an explicit Memberships'):
            read_workbook(self.path)

    def test_invalid_scores_and_statuses(self):
        for score, status, message in [(3, 'graded', 'exceeds question'), (None, 'graded', 'require a Score'),
                                       (0, 'absent', 'require a blank Score'), (None, 'unknown', 'unknown Status')]:
            with self.subTest(score=score, status=status):
                def change(book):
                    book['QuestionResults']['D2'] = score
                    book['QuestionResults']['E2'] = status
                self.edit(change)
                with self.assertRaisesRegex(WorkbookError, message):
                    read_workbook(self.path)

    def test_optional_tables_use_formula_validation(self):
        self.edit(lambda b: setattr(b['QuestionResults']['D2'], 'value', '=1+1'))
        with self.assertRaisesRegex(WorkbookError, 'formula has no saved value'):
            read_workbook(self.path)

    def test_incomplete_evidence_is_not_reconciled_or_filled_in(self):
        self.edit(lambda b: b['QuestionResults'].delete_rows(2))
        data = read_workbook(self.path)
        self.assertTrue(any('Reconciliation A1/1001: incomplete' in w for w in data.warnings))
        view = build_dashboard(data, '1001', assessment_id='A1')
        self.assertEqual(view.total_score, 3)
        self.assertEqual(view.topics[0].score, 1)

    def test_mismatched_totals_are_reported_and_retained(self):
        self.edit(lambda b: setattr(b['Results']['C2'], 'value', 4))
        data = read_workbook(self.path)
        self.assertTrue(any('question total 3/4 differs from Results 4/4' in w for w in data.warnings))
        self.assertEqual(build_dashboard(data, '1001', assessment_id='A1').total_score, 4)

    def test_invalid_assessment_boundaries(self):
        self.edit(lambda b: setattr(b['AssessmentBoundaries']['C3'], 'value', 0))
        with self.assertRaisesRegex(WorkbookError, 'strictly increasing'):
            read_workbook(self.path)

    def test_export_privacy_escaping_filters_and_freshness(self):
        def change(book):
            book['ExitTickets']['F2'] = '<script>unsafe-reflection</script>'
            book['ExitTickets']['F3'] = 'OTHER_STUDENT_REFLECTION'
            book['Questions']['F2'] = 'ANSWER_MUST_NOT_LEAK'
            book['Resources']['D2'] = '<img src=x onerror=alert(1)>'
        self.edit(change)
        generate_dashboards(self.config)
        html = (self.config.output / 'student1001.html').read_text(encoding='utf-8')
        self.assertNotIn('OTHER_STUDENT_REFLECTION', html)
        self.assertNotIn('ANSWER_MUST_NOT_LEAK', html)
        self.assertNotIn('student1002', html)
        self.assertIn('&lt;script&gt;unsafe-reflection&lt;/script&gt;', html)
        self.assertIn('&lt;img', html)
        self.assertIn('data-assessment="assessment-A2"', html)
        self.assertIn('id="assessment"', html)
        self.assertNotRegex(html, r'(?:src|href)=["\']https?://')
        client = create_app(self.config).test_client()
        self.assertNotIn(b'workbook has changed', client.get('/').data)
        self.edit(lambda b: setattr(b['ExitTickets']['G2'], 'value', 'New teacher feedback'))
        self.assertIn(b'workbook has changed', client.get('/').data)
        generate_dashboards(self.config)
        self.assertNotIn(b'workbook has changed', client.get('/').data)

    def test_duplicate_tickets_and_unknown_references(self):
        self.edit(lambda b: setattr(b['ExitTickets']['A3'], 'value', 'T1'))
        with self.assertRaisesRegex(WorkbookError, 'duplicate ticket'):
            read_workbook(self.path)
        self.edit(lambda b: setattr(b['ExitTickets']['B2'], 'value', 'unknown'))
        with self.assertRaisesRegex(WorkbookError, 'unknown assessment or student'):
            read_workbook(self.path)


if __name__ == '__main__':
    unittest.main()
