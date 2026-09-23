import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from hashlib import sha256

from openpyxl import Workbook, load_workbook
from create_example_workbook import create_example
from studentdash.classification import QuestionTag, classify, merge_tags, performance, syllabus_hierarchy
from studentdash.excel import WorkbookError, read_workbook
from tools.migrate_exam_databases import MANIFESTS, migrate
from tools.import_score_matrix import import_matrix


class ClassificationTests(unittest.TestCase):
    def test_validation_and_precedence(self):
        a = QuestionTag('Q', 'Skill', 'Chemical reasoning', 'teacher', 1)
        b = QuestionTag('Q', 'Skill', 'Data interpretation', 'teacher', None)
        c = QuestionTag('Q', 'Skill', 'Recall knowledge', 'rule', .8)
        self.assertEqual(set(merge_tags([a, a, b, c])), {a, b})
        existing = QuestionTag('Q', 'Skill', 'Recall knowledge', 'existing')
        self.assertEqual(merge_tags([existing, c]), [existing])
        for args in [('Skill', 'Unknown', 'rule', .5), ('Unknown', 'Recall', 'rule', .5),
                     ('CommandTerm', 'MCQ', 'existing', None), ('Skill', 'Recall knowledge', 'bad', None)]:
            with self.assertRaises(ValueError):
                QuestionTag('Q', *args)
        for value in [-1, 1.1, float('nan'), float('inf'), True, '0.5']:
            with self.assertRaises(ValueError):
                QuestionTag('Q', 'Skill', 'Recall knowledge', 'rule', value)

    def test_hierarchy_and_no_semantic_guessing(self):
        self.assertEqual(syllabus_hierarchy('S1.2.2'), ['S1.2.2', 'S1.2', 'S1'])
        self.assertEqual(syllabus_hierarchy('5.1'), ['5.1', '5'])
        self.assertEqual(syllabus_hierarchy('unknown'), [])
        tags = classify('Q', 'Explain this unfamiliar name. Calculate the gradient of the graph.', 'State and Explain')
        self.assertEqual({t.tag for t in tags if t.category == 'CommandTerm'}, {'State', 'Explain'})
        self.assertFalse(any(t.category in {'CognitiveDemand', 'Context'} for t in tags))
        self.assertIn('Graph', {t.tag for t in tags})
        quantitative = classify('Q', 'Calculate the mass of the product.')
        self.assertIn(('Skill', 'Quantitative problem solving'), {(t.category, t.tag) for t in quantitative})
        parts = classify('Q', '(a) Calculate. (b) Explain the graph.', 'Explain', ambiguous_parts=True)
        self.assertEqual([(t.category, t.tag) for t in parts], [('CommandTerm', 'Explain')])

    def test_migration_preserves_original_members_and_teacher_reviews(self):
        with tempfile.TemporaryDirectory() as folder:
            source, out = Path(folder)/'original.xlsx', Path(folder)/'classified.xlsx'
            book = Workbook()
            ws = book.active
            ws.append(['URL number', 'Question', 'Syllabus section', 'Command term', 'Paper', 'Level'])
            ws.append(['101', 'Calculate using the graph.', "['5.1 Measuring energy changes']", 'Calculate', '1', 'SL'])
            ws['B2'].number_format = '@'
            ws['G2'] = '=1+1'
            book.save(source)
            digest = sha256(source.read_bytes()).hexdigest()
            migrate(source, out)
            self.assertEqual(digest, sha256(source.read_bytes()).hexdigest())
            with ZipFile(source) as original, ZipFile(out) as upgraded:
                for name in original.namelist():
                    if name not in MANIFESTS:
                        self.assertEqual(original.read(name), upgraded.read(name), name)
            book = load_workbook(out)
            self.assertEqual(book.active['G2'].value, '=1+1')
            self.assertEqual(book['QuestionSyllabus']['D2'].value, 'unreviewed')
            self.assertIsNone(book['QuestionSyllabus']['C2'].value)
            book['QuestionTags'].append(['101', 'Representation', 'Written', 'teacher', 1])
            book.save(out)
            book.close()
            migrate(source, out)
            book = load_workbook(out)
            tags = [r for r in list(book['QuestionTags'].values)[1:] if r[1] == 'Representation']
            self.assertEqual([r[2] for r in tags], ['Written'])
            book.close()
            with self.assertRaises(ValueError):
                migrate(source, source)

    def test_workbook_tags_analytics_and_isolation(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'test.xlsx'
            create_example(path)
            book = load_workbook(path)
            ws = book['QuestionTags']
            for row in [['Q1', 'Skill', 'Chemical reasoning', 'teacher', 1],
                        ['Q1', 'Representation', 'Graph', 'teacher', 1],
                        ['Q1', 'Representation', 'Graph', 'teacher', 1]]:
                ws.append(row)
            ws = book['QuestionSyllabus']
            ws.append(['Q1', '2.1', 'S1.2.2', 'current', 'teacher'])
            book.save(path)
            book.close()
            data = read_workbook(path)
            filters = [('Skill', 'Chemical reasoning'), ('Representation', 'Graph')]
            for sid in ['1001', '1002']:
                result = performance(data, sid, filters, syllabus='S1')
                expected = [r for r in data.question_results if r.student_id == sid and r.question_id == 'Q1' and r.status == 'graded']
                self.assertEqual(result['score'], sum(r.score for r in expected))
                self.assertEqual(result['question_ids'], ['Q1'])
            self.assertEqual(performance(data, '1001', filters, syllabus='S2')['maximum'], 0)
            data.question_syllabus['Q1']['Status'] = 'out_of_scope'
            self.assertEqual(performance(data, '1001', filters, syllabus='S1')['maximum'], 0)
            with self.assertRaises(ValueError):
                performance(data, 'unknown')
            book = load_workbook(path)
            book['QuestionTags'].append(['Q1', 'Skill', 'bad', 'teacher', None])
            book.save(path)
            book.close()
            with self.assertRaises(WorkbookError):
                read_workbook(path)

    def test_matrix_import_and_rejection(self):
        with tempfile.TemporaryDirectory() as folder:
            source, out = Path(folder)/'test.xlsx', Path(folder)/'ready.xlsx'
            create_example(source)
            data = read_workbook(source)
            aid, sid = next(iter(data.memberships))
            from studentdash.question_data import eligible
            questions = [q for q in data.questions.values() if q.assessment_id == aid and eligible(q, data.memberships[aid, sid])]
            book = load_workbook(source)
            ws = book.create_sheet('ScoreMatrix')
            ws.append(['AssessmentID', 'StudentID']+[q.id for q in questions])
            ws.append([aid, sid]+[0 for q in questions])
            book.save(source)
            book.close()
            digest = sha256(source.read_bytes()).hexdigest()
            import_matrix(source, out)
            self.assertEqual(digest, sha256(source.read_bytes()).hexdigest())
            imported = read_workbook(out)
            self.assertEqual([r.score for r in imported.question_results if (r.assessment_id, r.student_id) == (aid, sid)], [0]*len(questions))
            with self.assertRaises(ValueError):
                import_matrix(source, out)
            book = load_workbook(source)
            book['ScoreMatrix']['C2'] = 999
            book.save(source)
            book.close()
            with self.assertRaises(WorkbookError):
                import_matrix(source, Path(folder)/'invalid.xlsx')
            self.assertFalse((Path(folder)/'invalid.xlsx').exists())
