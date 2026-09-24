from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from pypdf import PdfWriter

from studentdash.assessment_import import create_draft, draft_by_id, save_review, validation
from studentdash.classification import merge_tags
from studentdash.config import Config
from studentdash.document_extract import PDFImporter, DOCXImporter, ExtractedDocument, structure
from studentdash.entry import ClassStore, EntryError, as_workbook, save_assessment, save_scores, summary
from studentdash.import_suggestions import ProfileSuggestionProvider, curriculum_nodes, suggest_draft
from studentdash.render import generate_dashboards
from studentdash.teacher import create_app
from tools.create_import_example import docx_bytes, pdf_bytes


def review_payload(draft):
    return dict(title=draft['title'], date=draft['date'], description=draft['description'],
                participants=draft['participants'], review_ack=True, warning_note='',
                nodes=[dict(id=n['id'], number=n['number'], text=n['text'], marks=n['marks'], parent=n['parent'],
                            scorable=n['scorable'], tags=n.get('tags', {}), curriculum_nodes=n.get('curriculum_nodes', []), reviewed=True)
                       for n in draft['nodes']])


class ExtractionTests(unittest.TestCase):
    def test_pdf_docx_structure_marks_context_and_tables(self):
        for importer, content in [(PDFImporter(), pdf_bytes()), (DOCXImporter(), docx_bytes())]:
            with self.subTest(type=type(importer).__name__):
                parsed = structure(importer.extract(content))
                by_number = {n['number']: n for n in parsed['nodes']}
                self.assertEqual([n['number'] for n in parsed['nodes']], ['1', '2', '3', '3a', '3b', '4'])
                self.assertFalse(by_number['3']['scorable'])
                self.assertEqual(by_number['3b']['parent'], by_number['3']['id'])
                self.assertEqual(by_number['3b']['marks'], 2)
                self.assertIn('investigation', by_number['3']['text'])
                self.assertEqual(parsed['document_total'], 16)
                self.assertEqual(sum(n['marks'] for n in parsed['nodes'] if n['scorable']), 16)
                if isinstance(importer, PDFImporter):
                    self.assertEqual(by_number['1']['source_pages'], [1])
                else:
                    self.assertTrue(by_number['2']['tables'])
                    self.assertIn('20 | 4', by_number['2']['text'])

    def test_nested_roman_parts_and_inline_part(self):
        doc = ExtractedDocument(blocks=[dict(kind='text', text='1. (a) Use the evidence.\n(i) State a fact. [1]\n(ii) Explain a reason. [2]\n(b) Compare the samples. [3]', page=2)], method='test')
        nodes = structure(doc)['nodes']
        self.assertEqual([n['number'] for n in nodes], ['1', '1a', '1ai', '1aii', '1b'])
        self.assertEqual(nodes[2]['parent'], nodes[1]['id'])
        self.assertEqual(nodes[4]['parent'], nodes[0]['id'])

    def test_missing_marks_are_not_invented(self):
        nodes = structure(DOCXImporter().extract(docx_bytes(['1. Explain this relationship.'], False)))['nodes']
        self.assertIsNone(nodes[0]['marks'])
        self.assertEqual(nodes[0]['confidence']['marks'], 0)
        self.assertTrue(nodes[0]['warnings'])

    def test_conflicting_marks_require_manual_resolution(self):
        nodes = structure(DOCXImporter().extract(docx_bytes(['1. Explain this relationship. [2]', 'Give another reason. [3]'], False)))['nodes']
        self.assertIsNone(nodes[0]['marks'])
        self.assertEqual(nodes[0]['confidence']['marks'], 0)
        self.assertEqual(nodes[0]['mark_candidates'], [2, 3])

    def test_figure_reference_preserved(self):
        import base64
        doc = Document()
        doc.add_paragraph('1. Describe the figure. [2]')
        image = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')
        doc.add_picture(BytesIO(image))
        data = BytesIO()
        doc.save(data)
        node = structure(DOCXImporter().extract(data.getvalue()))['nodes'][0]
        self.assertTrue(node['figures'])
        self.assertTrue(any('Figures' in w for w in node['warnings']))


class AssessmentImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='studentdash-import-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = ClassStore(self.root / 'entered_classes')
        self.doc = self.store.create('Fictional class', 'Alice Private\nBob Private')
        self.key = self.doc['id']
        self.content = docx_bytes()
        self.did = create_draft(self.store, self.key, 'test.docx', self.content)

    def draft(self, did=None):
        return draft_by_id(self.store.read(self.key), did or self.did)

    def save(self, payload, finalize=False, did=None):
        return save_review(self.store, self.key, did or self.did, payload, self.store.read(self.key)['version'], finalize)

    def test_draft_separate_from_assessment_and_original_preserved(self):
        doc = self.store.read(self.key)
        self.assertEqual(doc['assessments'], [])
        self.assertEqual(self.draft()['status'], 'draft')
        filename, mime, content = self.store.source(self.key, self.did)
        self.assertEqual(filename, 'test.docx')
        self.assertEqual(content, self.content)
        self.assertIn('wordprocessingml', mime)
        reopened = ClassStore(self.root / 'entered_classes')
        self.assertEqual(reopened.source(self.key, self.did)[2], self.content)

    def test_automatic_suggestions_multiple_tags_and_confidence(self):
        draft = self.draft()
        first = draft['nodes'][0]
        selected = {(s['category'], s['value']) for s in first['suggestions']}
        self.assertIn(('CommandTerm', 'Calculate'), selected)
        self.assertIn(('Representation', 'Calculation'), selected)
        self.assertIn(('Skill', 'Quantitative problem solving'), selected)
        self.assertEqual(first['tags'], {})
        self.assertTrue(all(0 <= s['confidence'] <= 1 for n in draft['nodes'] for s in n['suggestions']))
        payload = review_payload(draft)
        payload['nodes'][0]['tags'] = {'Skill': ['Chemical reasoning', 'Quantitative problem solving'], 'Representation': ['Graph', 'Calculation']}
        self.save(payload)
        self.assertEqual(len(self.draft()['nodes'][0]['tags']['Skill']), 2)

    def test_save_draft_does_not_accept_or_reject_unreviewed_proposals(self):
        payload = review_payload(self.draft())
        for node in payload['nodes']:
            node['reviewed'] = False
        self.save(payload)
        self.assertTrue(all(s['decision'] == 'pending' for n in self.draft()['nodes'] for s in n['suggestions']))
        self.assertTrue(all(not n['tags'] for n in self.draft()['nodes']))

    def test_generic_curriculum_tree_definition_matching_and_no_invented_codes(self):
        profile = dict(name='Literature', categories={'Lens': ['Narrator', 'Audience']},
            tag_definitions={'Lens': {'Narrator': {'definition': 'Analyse the narrator perspective', 'keywords': ['narrator']}}},
            curriculum={'nodes': [dict(id='root', label='Reading', parent=None),
                dict(id='voice', label='Narrative voice', definition='Analyse narrator perspective', code='VOICE', parent='root', keywords=['narrator'])]})
        proposals = ProfileSuggestionProvider().suggest(dict(text='Analyse narrator perspective in this passage.'), '', profile)
        self.assertTrue(any(s['kind'] == 'curriculum' and s['value'] == 'voice' for s in proposals))
        self.assertTrue(any(s['category'] == 'Lens' for s in proposals))
        self.assertFalse(any(s['kind'] == 'curriculum' for n in self.draft()['nodes'] for s in n['suggestions']))
        profile['curriculum']['nodes'][0]['parent'] = 'voice'
        with self.assertRaises(EntryError):
            curriculum_nodes(profile)

    def test_teacher_overrides_and_rejections_survive_new_suggestions(self):
        payload = review_payload(self.draft())
        payload['nodes'][0]['tags'] = {'CommandTerm': ['Explain']}
        self.save(payload)
        draft = deepcopy(self.draft())
        suggest_draft(draft)
        first = draft['nodes'][0]
        self.assertEqual(first['tags'], {'CommandTerm': ['Explain']})
        self.assertEqual(next(s for s in first['suggestions'] if s['value'] == 'Calculate')['decision'], 'rejected')
        aid = self.save(review_payload(self.draft()), True)
        data = as_workbook(self.store.path(self.key))
        question = next(q for q in data.questions.values() if q.assessment_id == aid and q.number == '1')
        self.assertEqual({t.tag for t in merge_tags(data.question_tags) if t.question_id == question.id}, {'Explain'})
        self.assertTrue(all(t.source == 'teacher' for t in data.question_tags))

    def test_unknown_tags_invalid_confidence_and_unknown_curriculum(self):
        payload = review_payload(self.draft())
        payload['nodes'][0]['tags'] = {'Skill': ['Made up']}
        with self.assertRaisesRegex(EntryError, 'unknown classification'):
            self.save(payload)
        payload = review_payload(self.draft())
        payload['nodes'][0]['curriculum_nodes'] = ['invented-code']
        with self.assertRaises(EntryError):
            self.save(payload)
        class BadProvider:
            def suggest(self, *args):
                return [dict(kind='tag', category='Skill', value='Made up', source='inferred', confidence=.8),
                        dict(kind='tag', category='CommandTerm', value='Explain', source='rule', confidence=1.2)]
        draft = deepcopy(self.draft())
        suggest_draft(draft, BadProvider())
        self.assertTrue(all(not n['suggestions'] for n in draft['nodes']))

    def test_missing_marks_duplicate_labels_and_duplicate_content_block_finalization(self):
        for change, message in [('marks', 'maximum marks'), ('number', 'unique'), ('text', 'duplicate')]:
            payload = review_payload(self.draft())
            if change == 'marks':
                payload['nodes'][0]['marks'] = None
            elif change == 'number':
                payload['nodes'][1]['number'] = payload['nodes'][0]['number']
            else:
                payload['nodes'][1]['text'] = payload['nodes'][0]['text']
            with self.subTest(change=change), self.assertRaisesRegex(EntryError, message):
                self.save(payload, True)
        self.assertEqual(self.store.read(self.key)['assessments'], [])

    def test_mark_mismatch_requires_review_note_and_parent_structure_validated(self):
        payload = review_payload(self.draft())
        payload['nodes'][0]['marks'] = 1
        self.save(payload)
        self.assertTrue(validation(self.draft())['mismatch'])
        with self.assertRaisesRegex(EntryError, 'document-total difference'):
            self.save(payload, True)
        payload['warning_note'] = 'Checked the original; teacher changed question 1 to one mark.'
        self.assertIsNotNone(self.save(payload, True))

    def test_invalid_hierarchy_and_unreviewed_draft_cannot_finalize(self):
        payload = review_payload(self.draft())
        payload['nodes'][0]['parent'] = payload['nodes'][0]['id']
        with self.assertRaisesRegex(EntryError, 'cycle'):
            self.save(payload, True)
        payload = review_payload(self.draft())
        payload['nodes'][0]['reviewed'] = False
        with self.assertRaisesRegex(EntryError, 'Check each'):
            self.save(payload, True)
        payload['nodes'][0]['reviewed'] = True
        payload['review_ack'] = False
        with self.assertRaises(EntryError):
            self.save(payload, True)

    def test_failed_and_scanned_extraction_remain_recoverable_drafts(self):
        did = create_draft(self.store, self.key, 'broken.pdf', b'not a PDF')
        self.assertEqual(self.draft(did)['nodes'], [])
        self.assertTrue(self.draft(did)['extraction']['warnings'])
        self.assertEqual(self.store.source(self.key, did)[2], b'not a PDF')
        writer = PdfWriter()
        writer.add_blank_page(width=600, height=800)
        stream = BytesIO()
        writer.write(stream)
        scanned = create_draft(self.store, self.key, 'scan.pdf', stream.getvalue())
        self.assertTrue(any('no readable text' in w for w in self.draft(scanned)['extraction']['warnings']))

    def test_finalize_standard_scores_context_privacy_and_immutable_source(self):
        payload = review_payload(self.draft())
        aid = self.save(payload, True)
        doc = self.store.read(self.key)
        item = next(a for a in doc['assessments'] if a['id'] == aid)
        self.assertEqual(len(item['questions']), 5)
        q3b = next(q for q in item['questions'] if q['number'] == '3b')
        self.assertIn('investigation measures gas', q3b['text'])
        self.assertIn('Calculate the change', q3b['text'])
        self.assertTrue(q3b['import_context']['parent'])
        save_scores(self.store, self.key, aid, [['2', '3', '1', '2', '8'], ['0', 'P', 'A', 'E', 'M']], doc['version'])
        item = self.store.read(self.key)['assessments'][0]
        self.assertEqual(summary(item, doc['students'][0]['id']), (16, 16, 'graded'))
        save_assessment(self.store, self.key, aid, item, self.store.read(self.key)['version'])
        self.assertIn('import_context', self.store.read(self.key)['assessments'][0]['questions'][0])
        self.assertEqual(self.store.source(self.key, self.did)[2], self.content)
        config = Config(self.store.path(self.key), self.root / 'output')
        generate_dashboards(config, False)
        page = (config.output / f"student{doc['students'][0]['id']}.html").read_text(encoding='utf-8')
        for private in ['Bob Private', doc['students'][1]['id'], self.did, 'test.docx', 'Download original']:
            self.assertNotIn(private, page)
        with self.assertRaisesRegex(EntryError, 'already'):
            self.save(payload, True)

    def test_source_only_available_to_owning_class_and_teacher(self):
        app = create_app(Config(self.root / 'unused.xlsx', self.root / 'output'))
        app.testing = True
        client = app.test_client()
        client.get('/classes')
        route = f'/classes/{self.key}/imports/{self.did}/source'
        response = client.get(route)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, self.content)
        response.close()
        other = self.store.create('Other class', 'Other person')
        self.assertEqual(client.get(f'/classes/{other["id"]}/imports/{self.did}/source').status_code, 400)
        with client.session_transaction() as session:
            session['simulated_student'] = self.doc['students'][0]['id']
        self.assertEqual(client.get(route).status_code, 403)

    def test_http_upload_review_persistence_and_concurrent_rejection(self):
        app = create_app(Config(self.root / 'unused.xlsx', self.root / 'output'))
        app.testing = True
        client = app.test_client()
        client.get('/classes')
        with client.session_transaction() as session:
            token = session['csrf_token']
        response = client.post(f'/classes/{self.key}/import', data=dict(csrf_token=token, document=(BytesIO(pdf_bytes()), 'my-test.pdf')))
        self.assertEqual(response.status_code, 303)
        self.assertEqual(client.get(response.location).status_code, 200)
        url = f'/classes/{self.key}/imports/{self.did}'
        version = self.store.read(self.key)['version']
        payload = review_payload(self.draft())
        payload['nodes'][0]['marks'] = None
        response = client.post(url, data=dict(csrf_token=token, version=version, payload=json.dumps(payload), action='finalize'))
        self.assertEqual(response.status_code, 400)
        self.assertIn(b'positive maximum marks', response.data)
        self.save(payload)
        response = client.post(url, data=dict(csrf_token=token, version=version, payload=json.dumps(payload), action='save'))
        self.assertEqual(response.status_code, 400)
        self.assertIn(b'Another tab', response.data)
        self.assertEqual(client.post(f'/classes/{self.key}/import').status_code, 400)


if __name__ == '__main__':
    unittest.main()
