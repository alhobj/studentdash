"""Add classification sheets without reserializing any original worksheet.

Run from the project root: python tools/migrate_exam_databases.py
Outputs are replaced on rerun; reviewed tags/mappings in existing copies survive.
"""
import argparse
from collections import Counter
from copy import copy
from dataclasses import astuple
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
import sys
from tempfile import NamedTemporaryFile
import xml.etree.ElementTree as ET
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook, load_workbook
from studentdash.classification import VOCABULARY, RULES, QuestionTag, classify, merge_tags

MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
DOCREL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
ET.register_namespace('', MAIN)
ET.register_namespace('r', DOCREL)
MANIFESTS = {'xl/workbook.xml', 'xl/_rels/workbook.xml.rels', '[Content_Types].xml'}


def append_sheets(source, destination, sheets):
    """Copy original ZIP members byte-for-byte except the three sheet registries."""
    if source.resolve() == destination.resolve():
        raise ValueError('Output must not be the source workbook.')
    extra = Workbook()
    extra.remove(extra.active)
    for name, rows in sheets.items():
        ws = extra.create_sheet(name)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    buffer = BytesIO()
    extra.save(buffer)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(source) as original, ZipFile(buffer) as additions:
        workbook = ET.fromstring(original.read('xl/workbook.xml'))
        relationships = ET.fromstring(original.read('xl/_rels/workbook.xml.rels'))
        types = ET.fromstring(original.read('[Content_Types].xml'))
        listing = workbook.find(f'{{{MAIN}}}sheets')
        if set(sheets) & {s.attrib['name'] for s in listing}:
            raise ValueError('Source already contains classification sheets; use the original source.')
        sid = max(int(s.attrib['sheetId']) for s in listing)
        ids = {r.attrib['Id'] for r in relationships}
        payloads = {}
        sheet_nodes, relationship_nodes, type_nodes = [], [], []
        for i, name in enumerate(sheets, 1):
            rid = f'rIdClassification{i}'
            while rid in ids:
                rid += '_'
            ids.add(rid)
            target = f'worksheets/classification{i}.xml'
            if 'xl/' + target in original.namelist():
                raise ValueError('Classification ZIP member collision.')
            sheet_nodes.append(ET.tostring(ET.Element(f'{{{MAIN}}}sheet', {'name': name, 'sheetId': str(sid+i), f'{{{DOCREL}}}id': rid})))
            relationship_nodes.append(ET.tostring(ET.Element(f'{{{REL}}}Relationship', {'Id': rid, 'Type': DOCREL+'/worksheet', 'Target': target})))
            type_nodes.append(ET.tostring(ET.Element(f'{{{CT}}}Override', {'PartName': '/xl/'+target, 'ContentType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'})))
            payloads['xl/'+target] = additions.read(f'xl/worksheets/sheet{i}.xml')
        # Preserve namespace declarations, including prefixes referenced only by
        # mc:Ignorable; ElementTree reserialization would silently drop those.
        for member, closing, nodes in [
            ('xl/workbook.xml', b'</sheets>', sheet_nodes),
            ('xl/_rels/workbook.xml.rels', b'</Relationships>', relationship_nodes),
            ('[Content_Types].xml', b'</Types>', type_nodes),
        ]:
            content = original.read(member)
            if content.count(closing) != 1:
                raise ValueError(f'Unsupported XML registry layout: {member}')
            payloads[member] = content.replace(closing, b''.join(nodes)+closing)
        with NamedTemporaryFile(dir=destination.parent, suffix='.xlsx', delete=False) as temp:
            temporary = Path(temp.name)
        try:
            with ZipFile(temporary, 'w') as out:
                for info in original.infolist():
                    out.writestr(copy(info), payloads.pop(info.filename, original.read(info.filename)))
                for name, content in payloads.items():
                    out.writestr(name, content)
            with ZipFile(temporary) as verify:
                for name in original.namelist():
                    if name not in MANIFESTS and original.read(name) != verify.read(name):
                        raise ValueError(f'Original member changed: {name}')
            temporary.replace(destination)
        finally:
            temporary.unlink(missing_ok=True)


def inspect(source):
    book = load_workbook(source, read_only=True, data_only=False)
    records = []
    try:
        for sheet in book:
            rows = iter(sheet.values)
            headers = next(rows)
            if not {'URL number', 'Question', 'Syllabus section'} <= set(headers):
                continue
            for n, values in enumerate(rows, 2):
                row = dict(zip(headers, values))
                if row['URL number'] is None:
                    continue
                row['_sheet'], row['_row'] = sheet.title, n
                records.append(row)
    finally:
        book.close()
    ids = [str(r['URL number']) for r in records]
    if not records or len(ids) != len(set(ids)):
        raise ValueError(f'{source}: missing questions or duplicate source IDs.')
    return records


def prior_reviews(path):
    tags, syllabus = [], {}
    if not path.exists():
        return tags, syllabus
    book = load_workbook(path, read_only=True)
    try:
        if 'QuestionTags' in book:
            for row in list(book['QuestionTags'].values)[1:]:
                t = QuestionTag(*row[:5])
                if t.source == 'teacher':
                    tags.append(t)
        if 'QuestionSyllabus' in book:
            for row in list(book['QuestionSyllabus'].values)[1:]:
                if row[4] == 'teacher':
                    syllabus[str(row[0])] = list(row)
    finally:
        book.close()
    return tags, syllabus


def migrate(source, destination, records=None):
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError('Never overwrite the original.')
    original_hash = sha256(source.read_bytes()).hexdigest()
    records = inspect(source) if records is None else records
    tags, reviews = prior_reviews(destination)
    known_ids = {str(r['URL number']) for r in records}
    if any(t.question_id not in known_ids for t in tags) or set(reviews) - known_ids:
        raise ValueError('Reviewed IDs no longer exist in the source; refusing to discard reviews.')
    index = [['QuestionID', 'SourceSheet', 'SourceRow', 'ReferenceCode', 'Level', 'Paper', 'QuestionType', 'PartResolution', 'HasImage', 'SourceSyllabusCode']]
    syllabus = [['QuestionID', 'LegacyCode', 'CurrentCode', 'Status', 'Source']]
    issues = [['QuestionID', 'Issue']]
    for r in records:
        qid = str(r['URL number'])
        # Source exports repeat parent text for individual URL records. No reliable
        # URL-to-part join is present, so semantic tagging must abstain for these.
        ambiguous = bool(r.get('Question part'))
        tags.extend(classify(qid, r['Question'], r.get('Command term'), ambiguous_parts=ambiguous))
        codes = re.findall(r"(?:^|['\s])((?:\d+|[A-D])\.\d+(?:\.\d+)*)\b", str(r['Syllabus section']))
        legacy = codes[-1] if codes else ''
        syllabus.append(reviews.get(qid, [qid, legacy, '', 'unreviewed', 'existing']))
        image = '<img' in str(r['Question']).lower()
        index.append([qid, r['_sheet'], r['_row'], r.get('Reference code'), r.get('Level'), r.get('Paper'), 'MCQ' if str(r.get('Paper')) == '1' else '', 'ambiguous parent/part' if ambiguous else 'source row', image, legacy])
        if ambiguous:
            issues.append([qid, 'Part mapping unresolved: preserved row command terms; no text-derived tags.'])
        if image:
            issues.append([qid, 'Image content preserved, not visually classified; source HTML may be truncated.'])
        if not r.get('Command term'):
            issues.append([qid, 'No source command term.'])
    tags = merge_tags(tags)
    definitions = {
        'Direct': 'Retrieve taught knowledge or directly apply a taught procedure.',
        'Applied': 'Use known chemistry in a somewhat changed context.',
        'Unfamiliar': 'Determine how chemistry applies to unfamiliar information or context.',
        'Recall': 'Retrieve knowledge; independent of empirical difficulty.',
        'Apply': 'Use knowledge or procedures; judge the actual task.',
        'Analyse/Reason': 'Connect evidence and reason through relationships.',
        'Evaluate/Create': 'Judge alternatives or construct a justified approach.',
    }
    vocabulary = [['Category', 'Tag', 'Definition', 'Active']]
    for category, values in VOCABULARY.items():
        for tag in values:
            vocabulary.append([category, tag, definitions.get(tag, f'{category}: question requires or uses {tag.lower()}; review the actual prompt.'), True])
    notes = [
        ['Item', 'Explanation'],
        ['Original', source.name], ['SHA256', original_hash],
        ['QuestionID', 'Original URL number, stored as text; QuestionIndex maps it to the unchanged source row.'],
        ['Syllabus', 'Legacy codes are not current codes. All mappings start unreviewed; no legacy questions are automatically certified current.'],
        ['Status', 'unreviewed/current/partial/out_of_scope. Set Source=teacher after review; only current mappings enter current-syllabus queries.'],
        ['Teacher review', 'Edit QuestionTags with Source=teacher. Teacher category replaces automatic proposals on rerun. Reviews survive reruns.'],
        ['Unknowns', 'No automatic CognitiveDemand or Context. Marks, paper and command term alone cannot justify them.'],
        ['Evidence', 'Question text rules run only for resolvable rows. Specification, markscheme, examiner report, marks, level and paper remain available unchanged for review; no semantic inference from them.'],
        ['Confidence', 'Heuristic strength, not a calibrated probability; blank for existing metadata.'],
        ['Command extensions', 'Additional terms found in source metadata are retained as CommandTerm vocabulary; no new Skill or Representation values added.'],
        ['Part scope', 'Do not sum these database rows as assessment marks until parent/part mapping has been reviewed.'],
    ]
    notes.extend([f'{c}/{t}', f'Text regex: {pattern}; confidence={confidence}'] for c, t, pattern, confidence in RULES)
    append_sheets(source, destination, {'QuestionTags': [['QuestionID', 'Category', 'Tag', 'Source', 'Confidence']] + [list(astuple(t)) for t in tags], 'ClassificationVocabulary': vocabulary, 'QuestionIndex': index, 'QuestionSyllabus': syllabus, 'ClassificationReview': issues, 'ClassificationNotes': notes})
    if sha256(source.read_bytes()).hexdigest() != original_hash:
        raise ValueError('Original source changed during migration.')
    return records, tags, issues[1:]


def report(results, path):
    lines = ['# Classification report', '', 'Legacy syllabus database: current relevance is unreviewed for every automatically processed row. No current syllabus codes were invented.', '', 'Counts refer to source URL records, not independently verified question parts. Confidence is heuristic.', '']
    all_counts = Counter()
    for name, (rows, tags, issues) in results.items():
        lines += [f'## {name}', '', f'{len(rows)} source question records; {len(tags)} tags; {len(issues)} review flags.', '']
        counts = Counter((t.category, t.tag) for t in tags)
        all_counts.update(counts)
        for category in VOCABULARY:
            classified = len({t.question_id for t in tags if t.category == category})
            lines.append(f'- {category}: {classified}/{len(rows)} classified ({classified/len(rows):.1%}); unknown {(len(rows)-classified)/len(rows):.1%}.')
        covered = len({t.question_id for t in tags})
        lines += [f'- Any tag: {covered/len(rows):.1%}; no tags: {1-covered/len(rows):.1%}.', f'- Low confidence (<0.8): {sum(t.confidence is not None and t.confidence < .8 for t in tags)}.', '']
        for level in ('SL', 'HL'):
            ids = {str(r['URL number']) for r in rows if r.get('Level') == level}
            lines.append(f'{level}: {len(ids)} records; {sum(t.question_id in ids for t in tags)} tags.')
        lines += ['', '| Category | Tag | Count | SL | HL |', '|---|---|---:|---:|---:|']
        levels = {str(r['URL number']): r.get('Level') for r in rows}
        for category, values in VOCABULARY.items():
            for tag in values:
                selected = [t for t in tags if (t.category, t.tag) == (category, tag)]
                lines.append(f'| {category} | {tag} | {counts[category, tag]} | {sum(levels[t.question_id] == "SL" for t in selected)} | {sum(levels[t.question_id] == "HL" for t in selected)} |')
        lines += ['', 'Review flags (full ID-level list in ClassificationReview):']
        lines.extend(f'- {issue}: {n}' for issue, n in Counter(i[1] for i in issues).items())
        lines.append('')
    lines += ['## Findings and teacher decisions', '', f'Total source records: {sum(len(r[0]) for r in results.values())}.', '', '- Papers 2 and 3 repeat parent question text and lists of part labels. Their semantic tags are deliberately unknown until a teacher resolves the specific prompt for each URL record. This is a source granularity limitation, not evidence that those skills are absent.', '- Paper 1 has incomplete command metadata. Text matching can miss interrogative MCQs and can match terms in distractors; review all rule tags before assessment use.', '- Embedded images and sometimes truncated HTML limit text-only analysis. Image-based graphs, structures and spectra need visual review.', '- CognitiveDemand and Context remain entirely unknown automatically. Neither command terms nor an unfamiliar substance name establishes demand or transfer.', '- The supplied Skill vocabulary is retained. Recall, reasoning, data interpretation and multi-step solving need contextual review; zero automatic counts are not grounds for deleting them.', '- Experimental Uncertainty overlaps QuantitativeSkill Uncertainty deliberately: tag both when both are demonstrated. Keep error, accuracy/precision and evaluation distinct pending teacher review.', '- Chemical representation is broad; review examples before splitting it. No new representation or skill category is justified by this conservative pass.', '- Source command terms outside the examples (including Solve, Distinguish, List, Write, Label and Annotate) were retained; these are vocabulary additions for source preservation, not inferred new skills.', '- Legacy options do not establish current applicability. Review each question against your current syllabus; mark current, partial or out_of_scope, preserving the original legacy code. Partial questions require adaptation before inclusion in current-syllabus diagnostics.', '', '## Rule documentation', '', 'See ClassificationNotes in each workbook and studentdash/classification.py. Existing command terms take precedence; teacher tags replace automated tags in the same category. No difficulty labels are generated.']
    path.write_text('\n'.join(lines)+'\n', encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, default=Path('tests/past exams (old curriculum)'))
    parser.add_argument('--output-dir', type=Path, default=Path('output/exam_databases'))
    args = parser.parse_args()
    sources = [args.source_dir / f'Chem paper {i} all.xlsx' for i in (1, 2, 3)]
    inspected = [inspect(p) for p in sources]  # inspect all three before creating any copy
    results = {}
    for i, (source, rows) in enumerate(zip(sources, inspected), 1):
        destination = args.output_dir / f'Chem paper {i} classified.xlsx'
        results[source.name] = migrate(source, destination, rows)
        print(f'{destination}: {len(rows)} source records')
    report(results, args.output_dir / 'classification_report.md')


if __name__ == '__main__':
    main()
