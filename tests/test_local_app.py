"""Fictional regressions for the private assessment workflow. No real learner fixtures."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO, StringIO
import csv
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from app.analytics import dashboard
from app.classification_import import (headers, import_classifications, parse_file, preview,
                                       template, node_index)
from app.excel_import import import_results, read_grading_workbook, structural_changes
from app.generator import export_pages, generate, zip_pages
from app.storage import Store, ValidationError, curriculum_catalog, profile_catalog
from app.teacher import create_app


def workbook(dates=('25.09.2026', '09.10.2026'), names=('Fictional Alpha', 'Fictional Beta'),
             rows=None, labels=('1a', '1b'), maximum=(2, 3)):
    book = Workbook()
    overview = book.active
    overview.title = 'Overview'
    overview.append(['Name'])
    for name in names:
        overview.append([name])
    blueprint = book.create_sheet('Assessment blueprint')
    blueprint.append(['Task', *labels, 'Marks', 'x for absence', 'Percent', 'Suggested grade', 'Grade'])
    blueprint.append(['Marks', *maximum, '=SUM(B2:C2)'])
    for i in range(len(names)):
        blueprint.append([f'=Overview!A{i + 2}'])
    blueprint.append(['Sum'])
    blueprint.append(['Task'])
    for day in dates:
        sheet = book.copy_worksheet(blueprint)
        sheet.title = day
        for i, values in enumerate(rows if rows is not None else [(2, 1), (0, None)], 3):
            for j, value in enumerate(values, 2):
                sheet.cell(i, j, value)
    boundaries = book.create_sheet('Grade boundaries')
    boundaries.append(['Grade', 'Percent'])
    for grade, threshold in [(7, 80), (6, 65), (5, 54), (4, 44), (3, 32), (2, 16), (1, 0), ('FRAV', None)]:
        boundaries.append([grade, threshold])
    return save_book(book)


def save_book(book):
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def modify(content, callback):
    book = load_workbook(BytesIO(content))
    callback(book)
    return save_book(book)


def classification_file(key='Fictional_Class', assessment='2026-09-25', question='1a', **values):
    row = dict(ClassID=key, AssessmentID=assessment, QuestionID=question, **values)
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(row))
    writer.writeheader()
    writer.writerow(row)
    return parse_file(buffer.getvalue().encode(), 'classification.csv')


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = Store(self.root / 'private/studentdash.sqlite3')
        self.profile = profile_catalog()['ib_chemistry']
        self.curriculum = curriculum_catalog()['chemistry-2025']
        self.content = workbook()
        self.incoming = read_grading_workbook(self.content)
        self.key = 'Fictional_Class'

    def results(self, incoming=None, confirm=False, mappings=None):
        return import_results(self.store, self.key, incoming or self.incoming, self.store.revision(self.key),
                              self.profile, self.curriculum, confirm=confirm, mappings=mappings)

    def classified(self, **values):
        self.results()
        parsed = classification_file(**(values or dict(CommandTerm='Explain; Calculate', Skill='Chemical reasoning', SyllabusCode='S1.1')))
        return import_classifications(self.store, self.key, parsed, self.store.revision(self.key))

    def test_workbook_shape_and_formula_linked_roster(self):
        self.assertEqual(len(self.incoming['students']), 2)
        self.assertEqual(set(self.incoming['assessments']), {'2026-09-25', '2026-10-09'})
        assessment = self.incoming['assessments']['2026-09-25']
        self.assertEqual(set(assessment['questions']), {'1a', '1b'})
        self.assertEqual(assessment['results']['name:fictional beta']['1a'], dict(score=0, status='graded'))
        self.assertEqual(assessment['results']['name:fictional beta']['1b']['status'], 'pending')

    def test_blueprint_only_is_not_an_assessment(self):
        parsed = read_grading_workbook(workbook(dates=()))
        self.assertEqual(parsed['assessments'], {})
        self.assertTrue(any('No dated' in w for w in parsed['warnings']))

    def test_original_workbook_unchanged(self):
        path = self.root / 'input.xlsx'
        path.write_bytes(self.content)
        before = sha256(path.read_bytes()).hexdigest()
        self.results(read_grading_workbook(path.read_bytes()))
        self.assertEqual(sha256(path.read_bytes()).hexdigest(), before)

    def test_reimport_same_workbook_stable_ids_and_no_duplicates(self):
        first = self.results()
        second = self.results()
        self.assertEqual(first['students'], second['students'])
        self.assertEqual(first['assessments'], second['assessments'])
        self.assertEqual(len(second['assessments']), 2)

    def test_classifications_survive_restart_and_new_assessment(self):
        before = self.classified()
        changed = read_grading_workbook(workbook(dates=('25.09.2026', '09.10.2026', '23.10.2026')))
        after = self.results(changed)
        restarted = Store(self.store.path).read(self.key)
        self.assertEqual(len(restarted['assessments']), 3)
        self.assertEqual(before['assessments']['2026-09-25']['questions']['1a']['classification'], after['assessments']['2026-09-25']['questions']['1a']['classification'])
        self.assertNotIn('classification', restarted['assessments']['2026-10-23']['questions']['1a'])
        self.assertEqual(before['students'], after['students'])

    def test_reimport_classification_is_semantically_idempotent(self):
        doc = self.classified()
        parsed = classification_file(CommandTerm='Explain; Calculate', Skill='Chemical reasoning', SyllabusCode='S1.1')
        self.assertEqual(preview(doc, parsed)['changes'], [])
        updated = import_classifications(self.store, self.key, parsed, doc['revision'])
        self.assertEqual(doc['assessments'], updated['assessments'])

    def test_conflict_requires_explicit_review_and_audit(self):
        doc = self.classified()
        parsed = classification_file(CommandTerm='Evaluate')
        self.assertTrue(preview(doc, parsed)['conflicts'])
        with self.assertRaises(ValidationError):
            import_classifications(self.store, self.key, parsed, doc['revision'])
        updated = import_classifications(self.store, self.key, parsed, doc['revision'], replace=True)
        c = updated['assessments']['2026-09-25']['questions']['1a']['classification']
        self.assertEqual(c['tags']['CommandTerm'], ['Evaluate'])
        self.assertEqual(c['tags']['Skill'], ['Chemical reasoning'])
        self.assertIn('before', self.store.history(self.key)[0]['detail']['changes'][0])

    def test_blank_preserves_clear_removes_only_confirmed_field(self):
        doc = self.classified()
        blank = classification_file(CommandTerm='')
        self.assertFalse(preview(doc, blank)['changes'])
        clear = classification_file(CommandTerm='[clear]')
        with self.assertRaises(ValidationError):
            import_classifications(self.store, self.key, clear, doc['revision'])
        updated = import_classifications(self.store, self.key, clear, doc['revision'], replace=True)
        c = updated['assessments']['2026-09-25']['questions']['1a']['classification']
        self.assertEqual(c['tags']['CommandTerm'], [])
        self.assertEqual(c['tags']['Skill'], ['Chemical reasoning'])

    def test_unknown_ids_tags_and_duplicate_rows_are_atomic_errors(self):
        doc = self.results()
        for parsed in [classification_file(key='Other'), classification_file(assessment='2027-01-01'),
                       classification_file(question='unknown'), classification_file(Skill='Made up'),
                       classification_file(SyllabusCode='S1.1.999'), classification_file(Unconfigured='Value')]:
            with self.subTest(parsed=parsed):
                with self.assertRaises(ValidationError):
                    import_classifications(self.store, self.key, parsed, doc['revision'])
                self.assertEqual(self.store.read(self.key)['revision'], doc['revision'])
        parsed = classification_file(CommandTerm='State')
        parsed['rows'].append(parsed['rows'][0])
        with self.assertRaises(ValidationError):
            preview(doc, parsed)

    def test_partial_classification_does_not_guess_other_categories(self):
        doc = self.classified(CommandTerm='Explain')
        c = doc['assessments']['2026-09-25']['questions']['1a']['classification']
        self.assertEqual(c['tags'], {'CommandTerm': ['Explain']})
        view = dashboard(doc, '1001')
        self.assertEqual(view['groups']['Skill'], [])
        self.assertEqual(view['unknown']['Skill'], 4)

    def test_changed_marks_requires_confirmation_and_retires_classification(self):
        doc = self.classified()
        content = modify(self.content, lambda b: setattr(b['25.09.2026']['B2'], 'value', 4))
        incoming = read_grading_workbook(content)
        self.assertTrue(structural_changes(doc, incoming))
        with self.assertRaises(ValidationError):
            self.results(incoming)
        updated = self.results(incoming, confirm=True)
        q = updated['assessments']['2026-09-25']['questions']['1a']
        self.assertNotIn('classification', q)
        self.assertEqual(q['marks'], 4)
        self.assertEqual(q['retired'][0]['classification']['tags']['Skill'], ['Chemical reasoning'])

    def test_renamed_question_never_inherits_old_classification(self):
        self.classified()
        incoming = read_grading_workbook(modify(self.content, lambda b: setattr(b['25.09.2026']['B1'], 'value', '2a')))
        doc = self.results(incoming, confirm=True)
        questions = doc['assessments']['2026-09-25']['questions']
        self.assertFalse(questions['1a']['active'])
        self.assertIn('retired', questions['1a'])
        self.assertNotIn('classification', questions['2a'])
        doc = self.results(self.incoming, confirm=True)
        self.assertNotIn('classification', doc['assessments']['2026-09-25']['questions']['1a'])

    def test_removed_assessment_retains_history_but_not_current_analysis(self):
        self.classified()
        incoming = read_grading_workbook(workbook(dates=('09.10.2026',)))
        with self.assertRaises(ValidationError):
            self.results(incoming)
        doc = self.results(incoming, confirm=True)
        self.assertFalse(doc['assessments']['2026-09-25']['active'])
        self.assertIn('classification', doc['assessments']['2026-09-25']['questions']['1a'])
        self.assertEqual(len(dashboard(doc, '1001')['assessments']), 1)

    def test_student_reorder_and_confirmed_rename_keep_id(self):
        doc = self.results()
        reordered = read_grading_workbook(workbook(names=('Fictional Beta', 'Fictional Alpha')))
        doc = self.results(reordered)
        self.assertEqual(doc['students']['1001']['name'], 'Fictional Alpha')
        renamed = read_grading_workbook(workbook(names=('Fictional Renamed', 'Fictional Beta')))
        doc = self.results(renamed, mappings={'name:fictional renamed': '1001'})
        self.assertEqual(len(doc['students']), 2)
        self.assertEqual(doc['students']['1001']['name'], 'Fictional Renamed')

    def test_explicit_source_studentid_survives_name_changes(self):
        def add_ids(book):
            book['Overview']['B1'] = 'StudentID'
            book['Overview']['B2'] = 'external-alpha'
            book['Overview']['B3'] = 'external-beta'
        source = modify(self.content, add_ids)
        before = self.results(read_grading_workbook(source))
        renamed = modify(source, lambda b: setattr(b['Overview']['A2'], 'value', 'Updated fictional name'))
        after = self.results(read_grading_workbook(renamed))
        self.assertEqual(len(after['students']), 2)
        self.assertEqual(after['students']['1001']['name'], 'Updated fictional name')
        self.assertEqual(before['students']['1001']['source_key'], after['students']['1001']['source_key'])

    def test_ambiguous_identity_mapping_rolls_back_everything(self):
        before = self.results()
        renamed = read_grading_workbook(workbook(names=('Renamed One', 'Renamed Two')))
        with self.assertRaises(ValidationError):
            self.results(renamed, mappings={'name:renamed one': '1001', 'name:renamed two': '1001'})
        self.assertEqual(self.store.read(self.key), before)

    def test_ambiguous_names_and_duplicate_questions_rejected(self):
        with self.assertRaises(ValidationError):
            read_grading_workbook(workbook(names=('Same Name', ' SAME name ')))
        with self.assertRaises(ValidationError):
            read_grading_workbook(workbook(labels=('1a', '1a')))

    def test_statuses_zeros_and_absence_override(self):
        incoming = read_grading_workbook(workbook(rows=[('M', 'E'), (0, 'A')]))
        doc = self.results(incoming)
        a = dashboard(doc, '1001')['assessments'][0]
        self.assertEqual((a['missing'], a['exempt'], a['percent']), (1, 1, None))
        b = dashboard(doc, '1002')['assessments'][0]
        self.assertEqual((b['graded'], b['absent'], b['score']), (1, 1, 0))
        content = modify(self.content, lambda b: setattr(b['25.09.2026']['E3'], 'value', 'x'))
        incoming = read_grading_workbook(content)
        self.assertTrue(any('overrides' in w for w in incoming['warnings']))
        self.assertEqual({r['status'] for r in incoming['assessments']['2026-09-25']['results']['name:fictional alpha'].values()}, {'absent'})

    def test_invalid_scores_errors_and_uncached_score_formula(self):
        for value in (-1, 99, '#DIV/0!', '=SUM(1,1)', 'unknown'):
            content = modify(self.content, lambda b: setattr(b['25.09.2026']['B3'], 'value', value))
            with self.subTest(value=value), self.assertRaises(ValidationError):
                read_grading_workbook(content)

    def test_cached_question_formula_is_read_and_total_mismatch_warned(self):
        content = modify(self.content, lambda b: setattr(b['25.09.2026']['B3'], 'value', '=1+1'))
        # Fictional Excel cache, written into OOXML because openpyxl never calculates formulas.
        output = BytesIO()
        with ZipFile(BytesIO(content)) as source, ZipFile(output, 'w') as dest:
            for item in source.infolist():
                body = source.read(item.filename)
                if item.filename == 'xl/worksheets/sheet3.xml':
                    body = body.replace(b'<f>1+1</f><v></v>', b'<f>1+1</f><v>2</v>')
                dest.writestr(item, body)
        parsed = read_grading_workbook(output.getvalue())
        self.assertEqual(parsed['assessments']['2026-09-25']['results']['name:fictional alpha']['1a']['score'], 2)
        mismatch = read_grading_workbook(modify(self.content, lambda b: setattr(b['25.09.2026']['D3'], 'value', 999)))
        self.assertTrue(any('total differs' in w for w in mismatch['warnings']))

    def test_multi_tag_rollups_hierarchy_and_evidence_threshold(self):
        doc = self.classified(CommandTerm='Explain; Explain; Calculate', Skill='Chemical reasoning', SyllabusCode='S1; S1.1')
        a = dashboard(doc, '1001')
        terms = {r['tag']: r for r in a['groups']['CommandTerm']}
        self.assertEqual(terms['Explain']['score'], 2)
        self.assertEqual(terms['Calculate']['score'], 2)
        self.assertEqual(terms['Explain']['pattern'], 'Insufficient evidence')
        syllabus = {r['tag']: r for r in a['groups']['Syllabus']}
        s1 = next(v for k, v in syllabus.items() if k.startswith('S1 ·'))
        self.assertEqual((s1['score'], s1['maximum'], s1['questions']), (2, 2, 1))
        next_file = classification_file(assessment='2026-10-09', CommandTerm='Explain', Skill='Chemical reasoning')
        doc = import_classifications(self.store, self.key, next_file, doc['revision'])
        view = dashboard(doc, '1001')
        term = next(r for r in view['groups']['CommandTerm'] if r['tag'] == 'Explain')
        self.assertEqual((term['score'], term['maximum'], term['questions'], term['pattern']), (4, 4, 2, 'Strength'))
        other = dashboard(doc, '1002')
        term = next(r for r in other['groups']['CommandTerm'] if r['tag'] == 'Explain')
        self.assertEqual((term['score'], term['percent'], term['pattern']), (0, 0, 'Low performance'))

    def test_specific_syllabus_nodes_have_explicit_parents(self):
        doc = self.results()
        content = template(doc, scaffold=True)
        def add_nodes(book):
            book['SyllabusNodes'].append(['statement', 'chemistry-2025:S1.1', 'Fictional statement', 'S1.1.99'])
            sheet = book['Classifications']
            sheet.cell(2, 5, 'S1.1.99')
        parsed = parse_file(modify(content, add_nodes), 'classification.xlsx')
        doc = import_classifications(self.store, self.key, parsed, doc['revision'])
        self.assertIn('statement', node_index(doc['curriculum']))
        self.assertTrue(any(r['tag'].startswith('S1.1.99') for r in dashboard(doc, '1001')['groups']['Syllabus']))
        for nodes in [[dict(id='a', parent='b', label='A'), dict(id='b', parent='a', label='B')], [dict(id='a', parent='missing', label='A')]]:
            with self.assertRaises(ValidationError):
                node_index(dict(nodes=nodes))

    def test_templates_roundtrip_and_formula_injection_safe(self):
        doc = self.classified(Topic='=DANGEROUS()', Subtopic='<script>bad</script>')
        content = template(doc, scaffold=True)
        parsed = parse_file(content, 'template.xlsx')
        self.assertEqual(preview(doc, parsed)['changes'], [])
        book = load_workbook(BytesIO(content))
        self.assertNotEqual(book['Classifications']['F2'].data_type, 'f')
        book.close()
        formula = modify(content, lambda b: setattr(b['Classifications']['F2'], 'value', '=1+1'))
        with self.assertRaises(ValidationError):
            parse_file(formula, 'formula.xlsx')

    def test_different_subject_profile_requires_no_core_changes(self):
        profile = dict(name='Fictional music', categories={'Technique': ['Listening', 'Composition']})
        doc = import_results(self.store, 'Music', self.incoming, 0, profile, dict(nodes=[dict(id='root', parent=None, label='Music')]))
        parsed = classification_file(key='Music', Technique='Listening; Composition')
        doc = import_classifications(self.store, 'Music', parsed, doc['revision'])
        view = dashboard(doc, '1001')
        self.assertEqual({r['tag'] for r in view['groups']['Technique']}, {'Listening', 'Composition'})
        self.assertNotIn('Skill', view['groups'])

    def test_profile_level_scope_uses_explicit_metadata(self):
        doc = import_results(self.store, self.key, self.incoming, 0, self.profile, self.curriculum, default_level='SL')
        parsed = classification_file(Level='HL')
        doc = import_classifications(self.store, self.key, parsed, doc['revision'])
        self.assertEqual(len(dashboard(doc, '1001')['questions']), 3)
        parsed = classification_file(Level='BOTH')
        doc = import_classifications(self.store, self.key, parsed, doc['revision'], replace=True)
        self.assertEqual(len(dashboard(doc, '1001')['questions']), 4)

    def test_backup_restore_persistence_history_and_no_overwrite(self):
        doc = self.classified()
        backup = self.store.backup(self.root / 'backup.sqlite3')
        restored = Store.restore(backup, self.root / 'restored/studentdash.sqlite3')
        self.assertEqual(restored.read(self.key), doc)
        self.assertEqual(restored.history(self.key), self.store.history(self.key))
        with self.assertRaises(ValidationError):
            Store.restore(backup, restored.path)
        with self.assertRaises(ValidationError):
            self.store.backup(backup)

    def test_restore_rejects_invalid_schema_and_path_identifiers(self):
        self.results()
        backup = self.store.backup(self.root / 'invalid.sqlite3')
        with sqlite3.connect(backup) as db:
            row = db.execute('SELECT body FROM courses').fetchone()
            doc = json.loads(row[0])
            doc['students']['../escape'] = doc['students'].pop('1001')
            db.execute('UPDATE courses SET body=?', (json.dumps(doc),))
        destination = self.root / 'rejected.sqlite3'
        with self.assertRaises(ValidationError):
            Store.restore(backup, destination)
        self.assertFalse(destination.exists())

    def test_export_fingerprint_detects_different_database_at_same_revision(self):
        from app.generator import fingerprint
        doc = self.results()
        changed = deepcopy(doc)
        changed['students']['1001']['name'] = 'A restored alternate version'
        self.assertNotEqual(fingerprint(doc), fingerprint(changed))

    def test_stale_preview_rejected_without_partial_writes(self):
        doc = self.results()
        self.results()
        with self.assertRaises(ValidationError):
            import_classifications(self.store, self.key, classification_file(CommandTerm='State'), doc['revision'])
        self.assertNotIn('classification', self.store.read(self.key)['assessments']['2026-09-25']['questions']['1a'])

    def test_single_file_exports_private_clean_offline_and_regeneration(self):
        doc = self.classified(Topic='<script>bad</script>')
        pages = export_pages(doc)
        first, second = pages['student1001.html'], pages['student1002.html']
        self.assertNotIn('Fictional Beta', first)
        self.assertNotIn('Fictional Alpha', second)
        for html in pages.values():
            self.assertNotIn('<script', html)
            self.assertNotIn('src=', html)
            self.assertNotIn('href=', html)
            self.assertNotIn('studentdash.sqlite3', html)
        self.assertIn('&lt;script&gt;', first)
        directory = self.root / 'generated' / self.key
        generate(doc, directory)
        (directory / 'student9999.html').write_text('obsolete')
        generate(doc, directory)
        self.assertFalse((directory / 'student9999.html').exists())
        with ZipFile(BytesIO(zip_pages(doc))) as archive:
            self.assertEqual(set(archive.namelist()), {'student1001.html', 'student1002.html', 'READ-ME.txt'})
            self.assertNotIn(b'Fictional Beta', archive.read('student1001.html'))


class WebsiteTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.app = create_app(self.root / 'private', self.root / 'generated')
        self.app.testing = True
        self.client = self.app.test_client()
        self.client.get('/results')

    def token(self):
        with self.client.session_transaction() as session:
            return session['csrf']

    def preview_results(self, content=None):
        return self.client.post('/results', data=dict(csrf=self.token(), course_id='Fictional_Class',
                                file=(BytesIO(content or workbook()), 'grading.xlsx')), content_type='multipart/form-data')

    def import_results(self):
        response = self.preview_results()
        self.assertEqual(response.status_code, 302)
        review = self.client.get(response.location)
        self.assertIn(b'Nothing has been saved', review.data)
        self.assertEqual(self.app.extensions['store'].list(), [])
        imported = self.client.post(response.location, data=dict(csrf=self.token()))
        self.assertEqual(imported.status_code, 302)

    def test_full_website_upload_preview_confirm_classify_export_and_backup(self):
        self.import_results()
        for path in ('/students', '/classifications', '/dashboards', '/backup', '/student/1001'):
            self.assertEqual(self.client.get(path).status_code, 200, path)
        self.assertIn(b'Missing classifications', self.client.get('/students').data)
        for path in ('/classification-template.xlsx', '/classification-template.xlsx?questions=yes', '/classification-prompt.txt'):
            self.assertEqual(self.client.get(path).status_code, 200)
        parsed = classification_file(CommandTerm='Explain')
        body = 'ClassID,AssessmentID,QuestionID,CommandTerm\nFictional_Class,2026-09-25,1a,Explain\n'
        response = self.client.post('/classifications', data=dict(csrf=self.token(), file=(BytesIO(body.encode()), 'classifications.csv')))
        self.assertEqual(response.status_code, 302)
        self.assertIn(b'Proposed classifications', self.client.get(response.location).data)
        self.assertEqual(self.client.post(response.location, data=dict(csrf=self.token())).status_code, 302)
        self.assertIn(b'Partial classifications', self.client.get('/students').data)
        self.assertIn(b'stale', self.client.get('/dashboards').data)
        self.assertEqual(self.client.post('/dashboards', data=dict(csrf=self.token())).status_code, 302)
        with self.client.get('/generated/1001') as generated:
            self.assertEqual(generated.status_code, 200)
        self.assertIn(b'match the current', self.client.get('/dashboards').data)
        zipped = self.client.get('/dashboards.zip')
        self.assertEqual(zipped.status_code, 200)
        with ZipFile(BytesIO(zipped.data)) as archive:
            self.assertNotIn(b'Fictional Beta', archive.read('student1001.html'))
        with self.client.post('/backup', data=dict(csrf=self.token())) as backup:
            self.assertEqual(backup.status_code, 200)
            self.assertTrue(backup.data.startswith(b'SQLite format 3'))

    def test_post_csrf_host_checks_and_no_mutation_on_failed_upload(self):
        self.assertEqual(self.client.post('/results').status_code, 400)
        self.assertEqual(self.client.get('/results', headers={'Host': 'evil.invalid'}).status_code, 400)
        self.assertEqual(self.preview_results(b'invalid xlsx').status_code, 400)
        self.assertEqual(self.app.extensions['store'].list(), [])

    def test_rename_confirmation_is_required_and_preview_expires_on_revision(self):
        self.import_results()
        response = self.preview_results(workbook(names=('Renamed Alpha', 'Fictional Beta')))
        self.assertEqual(self.client.post(response.location, data=dict(csrf=self.token())).status_code, 400)
        confirmed = self.client.post(response.location, data={'csrf': self.token(), 'map:name:renamed alpha': '1001'})
        self.assertEqual(confirmed.status_code, 302)
        self.assertEqual(len(self.app.extensions['store'].read('Fictional_Class')['students']), 2)
        self.assertEqual(self.client.get(response.location).status_code, 400)

    def test_preview_survives_application_restart(self):
        response = self.preview_results()
        replacement = create_app(self.root / 'private', self.root / 'generated')
        replacement.testing = True
        other = replacement.test_client()
        with self.client.session_transaction() as source:
            state = dict(source)
        with other.session_transaction() as target:
            target.update(state)
        self.assertEqual(other.get(response.location).status_code, 200)
        self.assertEqual(other.post(response.location, data=dict(csrf=state['csrf'])).status_code, 302)
