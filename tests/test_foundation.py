import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from create_example_workbook import create_example
from studentdash.analytics import build_dashboard, suggested_grade
from studentdash.config import ROOT, Config
from studentdash.excel import SCHEMA, WorkbookError, read_workbook
from studentdash.render import generate_dashboards
from studentdash.teacher import create_app


def fixture(path):
    book = Workbook()
    book.remove(book.active)
    for name, headers in SCHEMA.items():
        book.create_sheet(name).append(list(headers) + (['Status'] if name == 'Results' else []))
    book['Students'].append(['1001', 'PRIVATE_ALICE', 'private-alice@example.invalid', 'Chemistry SL'])
    book['Students'].append(['1002', 'PRIVATE_BOB', 'private-bob@example.invalid', 'Chemistry SL'])
    book['Assessments'].append(['A1', 'Atomic structure', '2026-09-20', 'Chemistry SL', 10])
    book['Assessments'].append(['A2', 'Second assessment', '2026-09-21', 'Chemistry SL', 30])
    book['Questions'].append(['Q1', 'A1', '1', 'SL', 'State a particle.', 'PRIVATE_ANSWER', 10,
                              'Structure', 'Particles', 'Short', 'State', False])
    book['Results'].append(['A1', '1001', 8, 10, 'graded'])
    book['Results'].append(['A1', '1002', 3.14159, 10, 'graded'])
    for grade, percent in [(1, 0), (2, 16), (3, 32), (4, 44), (5, 54), (6, 65), (7, 80)]:
        book['Grade boundaries'].append([grade, percent])
    hidden = book.create_sheet('Private notes')
    hidden['A1'] = 'HIDDEN_PRIVATE_SENTINEL'
    hidden.sheet_state = 'hidden'
    book.save(path)
    book.close()


class FoundationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='studentdash-test-')
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.path = self.root / 'master.xlsx'
        self.config = Config(self.path, self.root / 'output')
        fixture(self.path)

    def edit(self, change):
        book = load_workbook(self.path)
        try:
            change(book)
            book.save(self.path)
        finally:
            book.close()

    def test_portable_example_workbook(self):
        example = self.root / 'example.xlsx'
        create_example(example)
        data = read_workbook(example)
        self.assertEqual(set(data.students), {'1001', '1002'})
        self.assertAlmostEqual(build_dashboard(data, '1001').overall_percent, 600 / 7)
        original = example.read_bytes()
        with self.assertRaises(FileExistsError):
            create_example(example)
        self.assertEqual(example.read_bytes(), original)

    def test_missing_file_sheet_and_column(self):
        with self.assertRaisesRegex(WorkbookError, 'Workbook not found'):
            read_workbook(self.root / 'missing.xlsx')
        self.edit(lambda b: b.remove(b['Questions']))
        with self.assertRaisesRegex(WorkbookError, 'Missing required sheet.*Questions'):
            read_workbook(self.path)
        fixture(self.path)
        self.edit(lambda b: setattr(b['Students']['A1'], 'value', 'WrongHeader'))
        with self.assertRaisesRegex(WorkbookError, 'Students row 1: missing column.*StudentID'):
            read_workbook(self.path)

    def test_formula_cache_and_excel_errors(self):
        self.edit(lambda b: setattr(b['Results']['C2'], 'value', '=4+4'))
        with self.assertRaisesRegex(WorkbookError, 'Results!C2.*formula has no saved value'):
            read_workbook(self.path)
        self.edit(lambda b: setattr(b['Results']['C2'], 'value', '#DIV/0!'))
        with self.assertRaisesRegex(WorkbookError, 'Results!C2.*Excel error'):
            read_workbook(self.path)

    def test_duplicate_and_unsafe_ids(self):
        self.edit(lambda b: setattr(b['Students']['A3'], 'value', '1001'))
        with self.assertRaisesRegex(WorkbookError, 'duplicate ID'):
            read_workbook(self.path)
        self.edit(lambda b: setattr(b['Students']['A3'], 'value', '../../other'))
        with self.assertRaisesRegex(WorkbookError, 'ID must contain'):
            read_workbook(self.path)

    def test_case_insensitive_collision(self):
        def change(b):
            b['Students']['A2'] = 'Abc'
            b['Students']['A3'] = 'abc'
        self.edit(change)
        with self.assertRaisesRegex(WorkbookError, 'colliding filenames'):
            read_workbook(self.path)

    def test_invalid_result_references_and_scores(self):
        self.edit(lambda b: setattr(b['Results']['A2'], 'value', 'unknown'))
        with self.assertRaisesRegex(WorkbookError, 'unknown AssessmentID'):
            read_workbook(self.path)
        fixture(self.path)
        self.edit(lambda b: setattr(b['Results']['C2'], 'value', 11))
        with self.assertRaisesRegex(WorkbookError, 'exceeds MaxScore'):
            read_workbook(self.path)
        self.edit(lambda b: setattr(b['Results']['C2'], 'value', -1))
        with self.assertRaisesRegex(WorkbookError, 'non-negative'):
            read_workbook(self.path)

    def test_duplicate_result(self):
        self.edit(lambda b: b['Results'].append(['A1', '1001', 8, 10, 'graded']))
        with self.assertRaisesRegex(WorkbookError, 'duplicate assessment/student'):
            read_workbook(self.path)

    def test_weighted_percent_and_grade_edges(self):
        self.edit(lambda b: b['Results'].append(['A2', '1001', 12, 30, 'graded']))
        data = read_workbook(self.path)
        view = build_dashboard(data, '1001')
        self.assertEqual(view.overall_percent, 50)
        self.assertEqual(view.total_score, 20)
        for percent, grade in [(0, 1), (15.99, 1), (16, 2), (32, 3), (44, 4), (54, 5), (65, 6), (80, 7), (100, 7)]:
            self.assertEqual(suggested_grade(percent, data.boundaries), grade)

    def test_statuses_and_empty_history(self):
        def change(b):
            b['Results']['C2'] = None
            b['Results']['E2'] = 'absent'
        self.edit(change)
        view = build_dashboard(read_workbook(self.path), '1001', False)
        self.assertIsNone(view.overall_percent)
        self.assertIsNone(view.history[0].grade)
        self.assertEqual(view.graded_count, 0)
        self.assertEqual(view.practice, [])
        self.edit(lambda b: setattr(b['Results']['E2'], 'value', 'graded'))
        with self.assertRaisesRegex(WorkbookError, 'graded results require'):
            read_workbook(self.path)
        self.edit(lambda b: b['Results'].delete_rows(2))
        self.assertEqual(build_dashboard(read_workbook(self.path), '1001').history, [])

    def test_invalid_boundaries(self):
        self.edit(lambda b: setattr(b['Grade boundaries']['B3'], 'value', 0))
        with self.assertRaisesRegex(WorkbookError, 'strictly increase'):
            read_workbook(self.path)

    def test_student_export_privacy_and_escaping(self):
        self.edit(lambda b: setattr(b['Assessments']['B2'], 'value', '<script>alert("x")</script>'))
        data, names = generate_dashboards(self.config)
        self.assertEqual(names, ['student1001.html', 'student1002.html'])
        for sid in data.students:
            html = (self.config.output / f'student{sid}.html').read_text(encoding='utf-8')
            for private in ['PRIVATE_ALICE', 'PRIVATE_BOB', 'private-alice@', 'private-bob@',
                            'PRIVATE_ANSWER', 'HIDDEN_PRIVATE_SENTINEL', '<script>']:
                self.assertNotIn(private, html)
            other = '1002' if sid == '1001' else '1001'
            self.assertNotIn(f'student{other}', html)
            self.assertIn('&lt;script&gt;', html)
            self.assertNotRegex(html, r'(?:src|href)=["\']https?://')
            self.assertIn('Fictional practice example', html)
        alice = (self.config.output / 'student1001.html').read_text(encoding='utf-8')
        self.assertNotIn('3.14159', alice)
        self.assertNotIn('31.4%', alice)

    def test_examples_do_not_change_real_results(self):
        data = read_workbook(self.path)
        example = build_dashboard(data, '1001')
        real_only = build_dashboard(data, '1001', False)
        self.assertEqual(example.history, real_only.history)
        self.assertEqual(example.overall_percent, real_only.overall_percent)
        self.assertEqual(example.strengths, ['Calculate'])
        self.assertEqual(example.work_on, ['Explain'])
        self.assertEqual(example.verbs[0].score, 2)
        self.assertEqual(example.verbs[0].maximum, 6)

    def test_generation_validation_preserves_previous_output(self):
        generate_dashboards(self.config)
        page = self.config.output / 'student1001.html'
        before = page.read_bytes()
        self.edit(lambda b: b.remove(b['Results']))
        with self.assertRaises(WorkbookError):
            generate_dashboards(self.config)
        self.assertEqual(page.read_bytes(), before)

    def test_teacher_generation_preview_and_request_protection(self):
        app = create_app(self.config)
        app.config['TESTING'] = True
        client = app.test_client()
        response = client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'master.xlsx', response.data)
        self.assertEqual(client.post('/generate').status_code, 400)
        with client.session_transaction() as session:
            token = session['csrf_token']
        response = client.post('/generate', data={'csrf_token': token}, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Generated 2 student dashboards', response.data)
        preview = client.get('/preview/student1001.html')
        self.assertEqual(preview.status_code, 200)
        self.assertNotIn(b'PRIVATE_BOB', preview.data)
        self.assertEqual(preview.headers['Cache-Control'], 'no-store')
        preview.close()
        self.assertEqual(client.get('/preview/master.xlsx').status_code, 404)
        self.assertEqual(client.get('/preview/student9999.html').status_code, 404)
        self.assertEqual(client.get('/', headers={'Host': 'foreign.example'}).status_code, 400)
        self.edit(lambda b: b.remove(b['Students']))
        self.assertIn(b'Missing required sheet', client.get('/').data)
        self.assertEqual(client.get('/preview/student1001.html').status_code, 404)

    def test_removed_students_are_not_served_and_stale_pages_removed(self):
        generate_dashboards(self.config)
        def change(b):
            b['Students'].delete_rows(3)
            b['Results'].delete_rows(3)
        self.edit(change)
        client = create_app(self.config).test_client()
        self.assertEqual(client.get('/preview/student1002.html').status_code, 404)
        generate_dashboards(self.config)
        self.assertFalse((self.config.output / 'student1002.html').exists())


if __name__ == '__main__':
    unittest.main()
