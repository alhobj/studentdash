"""Read the supplied schema without modifying Excel or guessing matrix joins."""
import math
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import Assessment, GradeBoundary, Question, Result, Student, WorkbookData

SCHEMA = {
    'Students': ('StudentID', 'StudentName', 'Email', 'Class'),
    'Assessments': ('AssessmentID', 'AssessmentName', 'Date', 'Subject', 'Marks'),
    'Questions': ('QuestionID', 'AssessmentID', 'QuestionNumber', 'SLHL', 'Text',
                  'Answer', 'Marks', 'Topic', 'Subtopic', 'QuestionType', 'ActionVerb', 'AutoMark'),
    'Results': ('AssessmentID', 'StudentID', 'Score', 'MaxScore'),
    'Grade boundaries': ('Grade', 'Percent'),
}
STATUSES = {'graded', 'missing', 'absent', 'exempt', 'pending'}
OPTIONAL_SCHEMA = {
    'QuestionTags': ('QuestionID', 'Category', 'Tag', 'Source', 'Confidence'),
    'QuestionSyllabus': ('QuestionID', 'LegacyCode', 'CurrentCode', 'Status', 'Source'),
    'RevisionAttempts': ('AttemptID', 'StudentID', 'AssessmentID', 'QuestionID', 'Date', 'Score', 'Note'),
    'Memberships': ('AssessmentID', 'StudentID', 'Level'),
    'QuestionResults': ('AssessmentID', 'QuestionID', 'StudentID', 'Score', 'Status'),
    'AssessmentBoundaries': ('AssessmentID', 'Grade', 'Percent'),
    'Resources': ('ResourceID', 'Topic', 'Title', 'Content'),
    'ExitTickets': ('TicketID', 'StudentID', 'AssessmentID', 'Date', 'Prompt', 'Response', 'Feedback'),
}


class WorkbookError(ValueError):
    """An actionable workbook error safe to show in the private teacher interface."""


def _text(value):
    return '' if value is None else str(value).strip()


def _id(value, location):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    value = _text(value)
    if not re.fullmatch(r'[A-Za-z0-9_-]+', value):
        raise WorkbookError(f'{location}: ID must contain letters, digits, underscores or hyphens.')
    return value


def _number(value, location, optional=False, positive=False):
    if value is None and optional:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise WorkbookError(f'{location}: expected a finite numeric value, found {value!r}.')
    if value < 0 or (positive and value == 0):
        raise WorkbookError(f'{location}: expected a {"positive" if positive else "non-negative"} number.')
    return float(value)


def _required(value, location):
    if not _text(value):
        raise WorkbookError(f'{location}: value is required.')
    return _text(value)


def read_workbook(path: Path) -> WorkbookData:
    path = Path(path)
    if not path.is_file():
        raise WorkbookError(f'Workbook not found: {path}. Run python create_example_workbook.py or set STUDENTDASH_WORKBOOK.')
    formulas = cached = None
    try:
        formulas = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        return _read(formulas, cached)
    except WorkbookError:
        raise
    except Exception as exc:
        raise WorkbookError(f'Could not read workbook {path.name}: {exc}') from exc
    finally:
        if formulas is not None:
            formulas.close()
        if cached is not None:
            cached.close()


def _read(formulas, cached):
    missing = set(SCHEMA) - set(formulas.sheetnames)
    if missing:
        raise WorkbookError('Missing required sheet(s): ' + ', '.join(sorted(missing)))
    data = WorkbookData()
    tables = {}
    used_formulas = False
    schema = {**SCHEMA, **{k: v for k, v in OPTIONAL_SCHEMA.items() if k in formulas.sheetnames}}
    for name, required in schema.items():
        source_rows = iter(formulas[name].iter_rows())
        value_rows = iter(cached[name].iter_rows())
        header_cells = next(source_rows)
        next(value_rows)
        headers = [_text(c.value) for c in header_cells]
        if len([h for h in headers if h]) != len(set(h for h in headers if h)):
            raise WorkbookError(f'{name} row 1: duplicate column headers.')
        missing_columns = set(required) - set(headers)
        if name == 'Results' and not {'Status', 'Stratus'} & set(headers):
            missing_columns.add('Status (or legacy Stratus)')
        if missing_columns:
            raise WorkbookError(f'{name} row 1: missing column(s): {", ".join(sorted(missing_columns))}.')
        if name == 'Results' and 'Status' in headers and 'Stratus' in headers:
            raise WorkbookError('Results: keep one status column, Status or Stratus, not both.')
        selected = set(required) | ({'Status', 'Stratus'} if name == 'Results' else set())
        rows = []
        for row_num, (source, values) in enumerate(zip(source_rows, value_rows), 2):
            if not any(c.value is not None for h, c in zip(headers, source) if h in selected):
                continue
            row = {}
            for col, (header, original, cell) in enumerate(zip(headers, source, values), 1):
                if header not in selected:
                    continue
                location = f'{name}!{get_column_letter(col)}{row_num} ({header})'
                if original.data_type == 'f':
                    used_formulas = True
                    if cell.value is None:
                        raise WorkbookError(f'{location}: formula has no saved value. Recalculate and save in Excel, then retry.')
                if cell.data_type == 'e':
                    raise WorkbookError(f'{location}: Excel error {cell.value}. Fix it and recalculate in Excel.')
                row[header] = cell.value
            rows.append((row_num, row))
        tables[name] = rows
    if used_formulas:
        data.warnings.append('Formula results use saved Excel caches; freshness cannot be verified. Recalculate and save in Excel before generating.')
    if 'Stratus' in [c.value for c in next(formulas['Results'].iter_rows())]:
        data.warnings.append('Results uses the legacy column Stratus; rename it to Status when updating the workbook.')

    def unique(target, key, record, location):
        if key in target:
            raise WorkbookError(f'{location}: duplicate ID {key}.')
        target[key] = record

    for n, r in tables['Students']:
        loc = f'Students row {n}'
        sid = _id(r['StudentID'], loc + ' StudentID')
        unique(data.students, sid, Student(sid, _required(r['StudentName'], loc + ' StudentName'),
                                         _text(r['Email']), _text(r['Class'])), loc)
    if not data.students:
        raise WorkbookError('Students: no student records found.')
    if len({s.casefold() for s in data.students}) != len(data.students):
        raise WorkbookError('Students: IDs produce colliding filenames on case-insensitive filesystems.')
    for n, r in tables['Assessments']:
        loc = f'Assessments row {n}'
        aid = _id(r['AssessmentID'], loc + ' AssessmentID')
        day = r['Date']
        if isinstance(day, datetime):
            day = day.date()
        if not isinstance(day, date):
            try:
                day = date.fromisoformat(str(day))
            except ValueError:
                raise WorkbookError(f'{loc} Date: use an Excel date or YYYY-MM-DD.') from None
        unique(data.assessments, aid, Assessment(aid, _required(r['AssessmentName'], loc + ' AssessmentName'),
               day, _required(r['Subject'], loc + ' Subject'), _number(r['Marks'], loc + ' Marks', positive=True)), loc)
    for n, r in tables['Questions']:
        loc = f'Questions row {n}'
        qid = _id(r['QuestionID'], loc + ' QuestionID')
        aid = _id(r['AssessmentID'], loc + ' AssessmentID')
        if aid not in data.assessments:
            raise WorkbookError(f'{loc}: unknown AssessmentID {aid}.')
        marks = _number(r['Marks'], loc + ' Marks', optional=True, positive=True)
        unique(data.questions, qid, Question(qid, aid, _required(r['QuestionNumber'], loc + ' QuestionNumber'),
               _text(r['SLHL']), _text(r['Text']), marks, _text(r['Topic']), _text(r['Subtopic']),
               _text(r['QuestionType']), _text(r['ActionVerb'])), loc)
        if marks is None or not all(_text(r[k]) for k in ('Text', 'Topic', 'Subtopic', 'ActionVerb')):
            data.warnings.append(f'{loc}: incomplete question marks/text/classification; real question analytics are unavailable.')
        if _text(r['ActionVerb']).upper() == 'MCQ':
            data.warnings.append(f'{loc}: MCQ is a question type; provide an IB command term in ActionVerb.')
    seen_results = set()
    for n, r in tables['Results']:
        loc = f'Results row {n}'
        aid, sid = _id(r['AssessmentID'], loc + ' AssessmentID'), _id(r['StudentID'], loc + ' StudentID')
        if aid not in data.assessments or sid not in data.students:
            raise WorkbookError(f'{loc}: unknown AssessmentID {aid} or StudentID {sid}.')
        if (aid, sid) in seen_results:
            raise WorkbookError(f'{loc}: duplicate assessment/student result ({aid}, {sid}).')
        seen_results.add((aid, sid))
        score = _number(r['Score'], loc + ' Score', optional=True)
        maximum = _number(r['MaxScore'], loc + ' MaxScore', positive=True)
        if score is not None and score > maximum:
            raise WorkbookError(f'{loc}: Score {score:g} exceeds MaxScore {maximum:g}.')
        status = _text(r.get('Status', r.get('Stratus'))).lower()
        if not status:
            status = 'graded' if score is not None else 'pending'
            data.warnings.append(f'{loc}: blank status interpreted as {status}.')
        if status not in STATUSES:
            raise WorkbookError(f'{loc}: unknown Status {status!r}; use {", ".join(sorted(STATUSES))}.')
        if (status == 'graded') != (score is not None):
            raise WorkbookError(f'{loc}: graded results require a Score; non-graded results require a blank Score.')
        assessment = data.assessments[aid]
        if maximum != assessment.marks and 'Memberships' not in tables:
            data.warnings.append(f'{loc}: MaxScore {maximum:g} differs from Assessments.Marks {assessment.marks:g}; using Results.MaxScore provisionally.')
        if 'Memberships' not in tables and data.students[sid].class_name != assessment.subject:
            data.warnings.append(f'{loc}: student Class differs from assessment Subject; confirm assessment membership.')
        data.results.append(Result(aid, sid, score, maximum, status))
    for n, r in tables['Grade boundaries']:
        loc = f'Grade boundaries row {n}'
        if _text(r['Grade']).upper() == 'FRAV' and r['Percent'] is None:
            continue
        grade = _number(r['Grade'], loc + ' Grade')
        percent = _number(r['Percent'], loc + ' Percent')
        if not grade.is_integer() or not 1 <= grade <= 7 or percent > 100:
            raise WorkbookError(f'{loc}: Grade must be an integer 1–7 and Percent must be 0–100.')
        data.boundaries.append(GradeBoundary(int(grade), percent))
    ordered = sorted(data.boundaries, key=lambda b: b.grade)
    if [b.grade for b in ordered] != list(range(1, 8)) or ordered[0].percent != 0:
        raise WorkbookError('Grade boundaries: require each grade 1–7 exactly once, with grade 1 starting at 0%.')
    if any(a.percent >= b.percent for a, b in zip(ordered, ordered[1:])):
        raise WorkbookError('Grade boundaries: minimum percentages must strictly increase with grade.')
    data.boundaries = ordered
    for aid, assessment in data.assessments.items():
        questions = [q for q in data.questions.values() if q.assessment_id == aid]
        if sum(q.marks or 0 for q in questions) != assessment.marks:
            data.warnings.append(f'Assessment {aid}: populated question marks do not match assessment Marks.')
    from .question_data import read_extensions
    read_extensions(data, tables)
    if 'QuestionResults' not in tables:
        data.warnings.append('QuestionResults is absent. Question drill-down is unavailable unless demonstration examples are enabled.')
    return data
