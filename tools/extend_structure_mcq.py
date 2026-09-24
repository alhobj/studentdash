"""Append 100 authored questions to a separate 200-question revision workbook.

Run: python tools/extend_structure_mcq.py
The content is isolated in tools/resources/structure_mcq_authored.json. Existing
questions, images, source sheets and input files are preserved. This is an artifact
builder for one chemistry collection, not a change to Studentdash's domain model.
"""
import argparse
from collections import Counter
from copy import copy
from hashlib import sha256
import json
from pathlib import Path
import posixpath
import random
import re
from tempfile import NamedTemporaryFile
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / 'output/exam_databases'
BASE = FOLDER / 'Structure 1.1-1.3 and 2.1-2.2 - 100 MCQ.xlsx'
OUTPUT = FOLDER / 'Structure 1.1-1.3 and 2.1-2.2 - 200 MCQ.xlsx'
CONTENT = ROOT / 'tools/resources/structure_mcq_authored.json'
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
ET.register_namespace('', NS)
ET.register_namespace('r', REL)


def tag(local):
    return f'{{{NS}}}{local}'


def read_tables(path, names):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        return {name: list(book[name].values) for name in names}
    finally:
        book.close()


def authored_questions(content=CONTENT):
    sections = json.loads(Path(content).read_text(encoding='utf-8'))['sections']
    if len(sections) != 5 or any(len(rows) != 20 for rows in sections.values()):
        raise ValueError('Require five sections of 20 authored questions.')
    rng = random.Random('Studentdash authored Structure collection v1')
    positions = list(range(4))*25
    rng.shuffle(positions)
    result = []
    for section, rows in sections.items():
        for index, row in enumerate(rows, 1):
            prompt, correct, wrong1, wrong2, wrong3, explanation, paper, source_id = row
            if not all(isinstance(s, str) and s.strip() for s in row[:6]):
                raise ValueError(f'{section}/{index}: missing content')
            if len({correct, wrong1, wrong2, wrong3}) != 4:
                raise ValueError(f'{section}/{index}: duplicate answer choice')
            options = [wrong1, wrong2, wrong3]
            rng.shuffle(options)
            position = positions[len(result)]
            options.insert(position, correct)
            qid = f'NEW_{section.replace(".", "")}_{index:03}'
            result.append(dict(id=qid, section=section, prompt=prompt,
                               text=prompt+'\n'+'\n'.join(f'{letter}. {option}' for letter, option in zip('ABCD', options)),
                               answer='ABCD'[position], correct=correct, explanation=explanation,
                               paper=paper, source_id=str(source_id)))
    if len({r['prompt'] for r in result}) != 100:
        raise ValueError('Duplicate new prompts.')
    assert Counter(r['answer'] for r in result) == dict.fromkeys('ABCD', 25)
    return result


def set_cell(row, col, value, style=None):
    address = f'{get_column_letter(col)}{row.attrib["r"]}'
    old = next((c for c in row if c.attrib.get('r') == address), None)
    if old is not None:
        if style is None:
            style = old.attrib.get('s')
        row.remove(old)
    attrs = {'r': address}
    if style is not None:
        attrs['s'] = style
    cell = ET.SubElement(row, tag('c'), attrs)
    if value is None:
        return
    if isinstance(value, (int, float)):
        ET.SubElement(cell, tag('v')).text = str(value)
    else:
        cell.set('t', 'inlineStr')
        text = ET.SubElement(ET.SubElement(cell, tag('is')), tag('t'))
        text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        text.text = str(value)
    # Keep cell references ordered for spreadsheet readers.
    row[:] = sorted(row, key=lambda c: (len(re.match(r'[A-Z]+', c.attrib['r'])[0]), c.attrib['r']))


def append_rows(xml, values, *, question_rows=False):
    data = xml.find(tag('sheetData'))
    templates = {int(row.attrib['r']): row for row in data}
    last = max(templates)
    columns = max(len(row) for row in values)
    for n, values_row in enumerate(values, last+1):
        template = templates.get(2 if n % 2 == 0 else 3, templates[last])
        attrs = {'r': str(n), 'ht': '60', 'customHeight': '1'}
        if question_rows:
            text = values_row[3]
            attrs['ht'] = str(min(409, max(125, (len(text)//95 + text.count('\n')+3)*15)))
        row = ET.SubElement(data, tag('row'), attrs)
        for col, value in enumerate(values_row, 1):
            original = next((c for c in template if re.match(r'[A-Z]+', c.attrib['r'])[0] == get_column_letter(col)), None)
            style = original.attrib.get('s') if original is not None else None
            set_cell(row, col, value, style)
    reference = f'A1:{get_column_letter(columns)}{last+len(values)}'
    xml.find(tag('dimension')).set('ref', reference)
    auto = xml.find(tag('autoFilter'))
    if auto is not None:
        auto.set('ref', reference)


def sheet_paths(archive):
    workbook = ET.fromstring(archive.read('xl/workbook.xml'))
    rels = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
    paths = {r.attrib['Id']: r.attrib['Target'] for r in rels}
    return {s.attrib['name']: (paths[s.attrib[f'{{{REL}}}id']].lstrip('/')
            if paths[s.attrib[f'{{{REL}}}id']].startswith('/') else
            posixpath.normpath('xl/'+paths[s.attrib[f'{{{REL}}}id']]))
            for s in workbook.find(tag('sheets'))}


def build(base=BASE, output=OUTPUT, content=CONTENT):
    base, output = Path(base), Path(output)
    if output.exists() or base.resolve() == output.resolve():
        raise FileExistsError('Choose a new output path; existing workbooks are never overwritten.')
    new = authored_questions(content)
    names = ['Read me', 'Questions', 'AnswerKey', 'MappingReview', 'QuestionTags']
    original = read_tables(base, names)
    if len(original['Questions']) != 101 or len(original['AnswerKey']) != 101:
        raise ValueError('Expected the original 100-question collection.')
    old_ids = {r[2] for r in original['Questions'][1:]}
    old_text = {re.sub(r'\s+', ' ', r[3]).casefold() for r in original['Questions'][1:]}
    for r in new:
        if r['id'] in old_ids or re.sub(r'\s+', ' ', r['text']).casefold() in old_text:
            raise ValueError('New question duplicates an existing record.')
    input_paths = [base]+[FOLDER / f'Chem paper {p} classified.xlsx' for p in (1, 2, 3)]
    hashes = {p: sha256(p.read_bytes()).hexdigest() for p in input_paths}
    source_rows = {}
    for paper, path in enumerate(input_paths[1:], 1):
        table = read_tables(path, ['Sheet1'])['Sheet1']
        source_rows[paper] = {str(row[0]): dict(zip(table[0], row)) for row in table[1:]}
    questions, answers, mappings = [], [], []
    for n, r in enumerate(new, 101):
        src = source_rows[r['paper']][r['source_id']]
        questions.append([n, r['section'], r['id'], r['text'], '', 'SL / HL core', 'Newly authored', None, None, None, None, 1])
        answers.append([n, r['id'], r['section'], r['answer'], r['explanation'], 'Newly authored'])
        mappings.append([r['id'], r['section'], 'assistant authorship and content review',
                         'proposed — teacher confirmation pending',
                         'New wording and answer choices; the source is topic inspiration, not the origin of this question.',
                         r['paper'], r['source_id'], src['Reference code'], 'Topic inspiration only'])
    with ZipFile(base) as source:
        paths = sheet_paths(source)
        edited = {name: ET.fromstring(source.read(paths[name])) for name in names if name != 'QuestionTags'}
        append_rows(edited['Questions'], questions, question_rows=True)
        append_rows(edited['AnswerKey'], answers)
        data = edited['AnswerKey'].find(tag('sheetData'))
        set_cell(data[0], 5, 'Explanation for adapted / newly authored question')
        append_rows(edited['MappingReview'], mappings)
        header = edited['MappingReview'].find(tag('sheetData'))[0]
        for col, value in enumerate(['InspirationPaper', 'InspirationSourceID', 'InspirationReference', 'Relationship'], 6):
            set_cell(header, col, value, header[0].attrib.get('s'))
        cols = edited['MappingReview'].find(tag('cols'))
        for index, width in [(6, 20), (7, 24), (8, 28), (9, 30)]:
            ET.SubElement(cols, tag('col'), {'min': str(index), 'max': str(index), 'width': str(width), 'customWidth': '1'})
        intro = edited['Read me'].find(tag('sheetData'))
        set_cell(intro[0], 1, '200 MCQs — Structure 1.1, 1.2, 1.3, 2.1 and 2.2')
        by_section = Counter(r[1] for r in original['Questions'][1:])
        by_section.update(r['section'] for r in new)
        replacements = {
            'Contents': 'Questions 1–100: 86 original Paper 1 MCQs and 14 adaptations. Questions 101–200: 100 newly authored MCQs, 20 per section. All carry one mark.',
            'Answers': 'Original answers retained. Every new question has one keyed answer and an explanation. Newly authored questions are not official IB questions.',
            'Classification': 'Existing QuestionTags retained unchanged. New questions have section mappings but no diagnostic tags assigned. MappingReview distinguishes new authorship and topic inspiration from source originals.',
            'Coverage': 'New questions: 20 each in S1.1, S1.2, S1.3, S2.1 and S2.2. Combined totals appear below; coverage of every syllabus statement is not claimed.',
            'Use': 'Filter Type = Newly authored for questions 101–200. AnswerKey is separate. Inspiration references for new questions are in MappingReview; their Source paper/ID fields are blank because they are not past-paper questions.',
        }
        for index, values in enumerate(original['Read me']):
            key = values[0]
            if key in replacements:
                set_cell(intro[index], 2, replacements[key])
            elif key and key.split(' ')[0] in by_section:
                set_cell(intro[index], 2, by_section[key.split(' ')[0]])
        append_rows(edited['Read me'], [
            ['New questions', 'Written by the assistant on 2026-09-24 using the same three banks for topic coverage and style. Correct answers and explanations were written independently; no claim of IB authorship or teacher confirmation.'],
            ['Answer distribution', 'For the 100 newly authored questions, A, B, C and D each occur 25 times. Options are shuffled reproducibly; the question order is by section.'],
            ['Preservation', 'The original 100-question workbook and the three classified banks were not modified. Original 100 question rows, answer values, 18 embedded images and source records remain in this copy.'],
            ['Original collection SHA256', hashes[base]],
        ])
        modified = {paths[name]: ET.tostring(xml, encoding='utf-8') for name, xml in edited.items()}
        output.parent.mkdir(parents=True, exist_ok=True)
        with NamedTemporaryFile(dir=output.parent, suffix='.xlsx', delete=False) as handle:
            temporary = Path(handle.name)
        try:
            with ZipFile(temporary, 'w') as dest:
                for info in source.infolist():
                    dest.writestr(copy(info), modified.get(info.filename, source.read(info.filename)))
            validate(temporary, base, original, new, set(modified))
            if any(sha256(p.read_bytes()).hexdigest() != digest for p, digest in hashes.items()):
                raise ValueError('An input changed during generation.')
            temporary.replace(output)
        finally:
            temporary.unlink(missing_ok=True)
    print(output)
    print('200 questions; new questions 101–200; 20 new per section; 25 of each answer letter.')
    print('All original question/answer values, source records and embedded image bytes verified.')


def validate(output, base, original, new, modified):
    tables = read_tables(output, ['Questions', 'AnswerKey', 'MappingReview', 'QuestionTags'])
    if tables['Questions'][:101] != original['Questions']:
        raise ValueError('An original question changed.')
    if tables['AnswerKey'][1:101] != original['AnswerKey'][1:]:
        raise ValueError('An original answer changed.')
    if tables['QuestionTags'] != original['QuestionTags']:
        raise ValueError('Original tags changed.')
    assert len(tables['Questions']) == len(tables['AnswerKey']) == 201
    assert len({r[2] for r in tables['Questions'][1:]}) == 200
    for question, answer, expected in zip(tables['Questions'][101:], tables['AnswerKey'][101:], new):
        assert question[2] == answer[1] == expected['id']
        assert question[3] == expected['text']
        selected = next(line[3:] for line in question[3].splitlines() if line.startswith(answer[3]+'. '))
        assert selected == expected['correct']
        assert answer[4] == expected['explanation'] and answer[4]
        assert question[7:11] == (None, None, None, None)
    with ZipFile(base) as a, ZipFile(output) as b:
        assert b.testzip() is None
        for name in a.namelist():
            if name not in modified:
                assert a.read(name) == b.read(name), name
        assert len([n for n in b.namelist() if n.startswith('xl/media/')]) == 18


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base', type=Path, default=BASE)
    parser.add_argument('--output', type=Path, default=OUTPUT)
    parser.add_argument('--content', type=Path, default=CONTENT)
    args = parser.parse_args()
    build(args.base, args.output, args.content)
