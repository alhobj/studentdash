"""Document adapters produce the same source blocks; structure extraction is separate.

This deliberately handles simple numbered documents, not arbitrary exam layouts.
Source text, tables and image references remain available even if recognition fails.
"""
from dataclasses import dataclass, field
from io import BytesIO
import re
from typing import Protocol
from zipfile import ZipFile

from .entry import identifier

MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 200_000


class ExtractionError(ValueError):
    pass


@dataclass
class ExtractedDocument:
    blocks: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    method: str = ''
    metadata: dict = field(default_factory=dict)


class AssessmentImporter(Protocol):
    def extract(self, content: bytes) -> ExtractedDocument: ...


class PDFImporter:
    def extract(self, content):
        from pypdf import PdfReader
        result = ExtractedDocument(method='pdf-text-v1')
        try:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ExtractionError('This PDF is password protected. Upload an unlocked copy.')
            if len(reader.pages) > 50:
                raise ExtractionError('Use a document of at most 50 pages.')
            result.metadata['pages'] = len(reader.pages)
            for page_number, page in enumerate(reader.pages, 1):
                source = page.extract_text() or ''
                resources = page.get('/Resources', {})
                resources = resources.get_object() if hasattr(resources, 'get_object') else resources
                objects = resources.get('/XObject', {})
                objects = objects.get_object() if hasattr(objects, 'get_object') else objects
                images = [f'Page {page_number}: image {name}' for name, value in objects.items()
                          if value.get_object().get('/Subtype') == '/Image']
                if not source.strip():
                    result.warnings.append(f'Page {page_number}: no readable text. Scanned pages need manual transcription; OCR is not available yet.')
                result.blocks.append(dict(kind='text', text=source, page=page_number, figures=images))
                if sum(len(b['text']) for b in result.blocks) > MAX_TEXT:
                    raise ExtractionError('This document contains too much text. Upload a shorter test.')
            result.warnings.append('PDF reading order, equations, tables and figure associations need comparison with the original.')
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError('The PDF could not be read. Export a fresh PDF or try DOCX.') from exc
        return result


class DOCXImporter:
    def extract(self, content):
        from docx import Document
        from docx.table import Table
        try:
            with ZipFile(BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024:
                    raise ExtractionError('The DOCX expands to too much content. Upload a smaller test.')
                if 'word/document.xml' not in archive.namelist():
                    raise ExtractionError('This file is not a valid DOCX document.')
            doc = Document(BytesIO(content))
            result = ExtractedDocument(method='docx-blocks-v1', metadata={'title': doc.core_properties.title or ''})
            for index, item in enumerate(doc.iter_inner_content(), 1):
                if isinstance(item, Table):
                    cells = [[c.text for c in row.cells] for row in item.rows]
                    result.blocks.append(dict(kind='table', text='\n'.join(' | '.join(row) for row in cells),
                                              cells=cells, page=None, location=f'Table {index}', figures=[]))
                else:
                    figures = []
                    for drawing in item._p.xpath('.//wp:docPr'):
                        figures.append(drawing.get('descr') or drawing.get('name') or f'Figure in paragraph {index}')
                    result.blocks.append(dict(kind='text', text=item.text, page=None,
                                              location=f'Paragraph {index}', figures=figures))
                    if item._p.xpath('.//w:numPr'):
                        result.warnings.append('Automatic Word numbering was detected. Check question labels; use printed labels if they were not extracted.')
                    if item._p.xpath('.//m:oMath'):
                        result.warnings.append('Word equations require comparison with the original; equation layout is not extracted.')
            result.warnings.append('DOCX page numbers are unavailable. Paragraph/table references identify source locations.')
            if sum(len(b['text']) for b in result.blocks) > MAX_TEXT:
                raise ExtractionError('This document contains too much text. Upload a shorter test.')
            return result
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError('The DOCX could not be read. Save a fresh .docx copy and retry.') from exc


IMPORTERS = {'.pdf': PDFImporter, '.docx': DOCXImporter}
ROOT_QUESTION = re.compile(r'^(?:Question\s+|Q\s*)?(\d{1,3})(?:[.)]\s*|\s+|(?=\([a-z]+\))|$)(.*)$', re.I)
PART = re.compile(r'^\(([a-z]{1,4})\)\s*(.*)$', re.I)
MARK = re.compile(r'(?:\[(\d+(?:\.\d+)?)\]|\((\d+(?:\.\d+)?)\s*marks?\)|(\d+(?:\.\d+)?)\s+marks?)\s*$', re.I)
TOTAL = re.compile(r'^(?:Total(?:\s+maximum)?(?:\s+marks)?|Maximum\s+marks)\s*[:=]?\s*(\d+(?:\.\d+)?)\s*(?:marks?)?\s*$', re.I)


def structure(document):
    """Recognize root/letter/roman headings and trailing explicit mark annotations."""
    nodes, preamble, totals = [], [], []
    current = root = letter = None

    def node(number, parent, block):
        value = dict(id=identifier(), number=number, parent=parent['id'] if parent else None,
                     text='', marks=None, scorable=True, source_text='', source_pages=[], source_locations=[],
                     figures=[], tables=[], confidence={'number': .95, 'text': .75 if document.method.startswith('pdf') else .9, 'marks': 0},
                     warnings=[], reviewed=False, tags={}, curriculum_nodes=[], decisions={}, reviewed_categories=[])
        nodes.append(value)
        if parent:
            parent['scorable'] = False
        return value

    def append(value, raw, block):
        value['source_text'] += ('\n' if value['source_text'] else '') + raw
        for field, source_key in [('source_pages', 'page'), ('source_locations', 'location')]:
            if block.get(source_key) is not None and block[source_key] not in value[field]:
                value[field].append(block[source_key])
        value['figures'] = list(dict.fromkeys(value['figures'] + block.get('figures', [])))
        mark = MARK.search(raw)
        display = raw
        if mark:
            maximum = float(next(v for v in mark.groups() if v is not None))
            candidates = value.setdefault('mark_candidates', [])
            candidates.append(maximum)
            if len(candidates) > 1 or re.search(r'\[\d+(?:\.\d+)?\]', raw[:mark.start()]):
                value['warnings'].append('More than one mark annotation found. Check whether this question needs splitting.')
                value['marks'] = None
                value['confidence']['marks'] = 0
            else:
                value['marks'] = maximum
                value['confidence']['marks'] = .95
            display = raw[:mark.start()].strip()
        if display:
            value['text'] += ('\n' if value['text'] else '') + display

    for block in document.blocks:
        if block['kind'] == 'table':
            if current:
                current['tables'].append(block)
                current['text'] += '\n' + block['text']
                current['source_text'] += '\n' + block['text']
                current['warnings'].append('Check that this table belongs to this question or its parent context.')
            else:
                preamble.append(block['text'])
            continue
        for raw in block['text'].splitlines() or ([''] if block.get('figures') else []):
            line = raw.strip()
            total = TOTAL.match(line)
            if total:
                totals.append(float(total.group(1)))
                continue
            if not line and not block.get('figures'):
                continue
            top = ROOT_QUESTION.match(line)
            if top:
                root = current = node(top.group(1), None, block)
                letter = None
                line = top.group(2)
            part = PART.match(line)
            if part and root:
                label, line = part.groups()
                label = label.lower()
                parent = letter if letter and re.fullmatch(r'(?:i|ii|iii|iv|v|vi|vii|viii|ix|x)', label) else root
                current = node(parent['number'] + label, parent, block)
                if parent is root:
                    letter = current
            if current:
                append(current, line, block)
            else:
                preamble.append(raw)
    for value in nodes:
        if value['scorable'] and value['marks'] is None:
            value['warnings'].append('Maximum marks were not found. Enter the marks from the source.')
        if value['scorable'] and not value['text'].strip():
            value['warnings'].append('Question text is missing. Check the source and enter the prompt.')
        if value['figures']:
            value['warnings'].append('Figures remain in the original document. Check their association and any needed visual context.')
        if value['warnings']:
            value['confidence']['text'] = min(value['confidence']['text'], .6)
    warnings = list(dict.fromkeys(document.warnings))
    if not nodes:
        warnings.append('No numbered questions were reliably detected. The source is preserved; add or transcribe questions in review.')
    if len(set(totals)) > 1:
        warnings.append('Several different document totals were found. Confirm the intended total.')
    return dict(title=document.metadata.get('title') or next((s.strip() for s in preamble if s.strip()), ''),
                instructions='\n'.join(preamble), document_total=totals[0] if len(set(totals)) == 1 else None,
                detected_totals=totals, nodes=nodes, warnings=warnings, method=document.method,
                metadata=document.metadata, blocks=document.blocks)
