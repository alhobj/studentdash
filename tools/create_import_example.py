"""Build small fictional PDF/DOCX examples of the supported numbered layout."""
from io import BytesIO
from pathlib import Path


LINES = [
    'Fictional assessment import example',
    'Total marks: 16',
    'Answer all questions. Show your working.',
    '1. Calculate the mass of 0.5 mol of a substance with molar mass 20 g/mol. [2]',
    '2. Use the data to explain why the reaction rate changes. [3]',
    'Temperature (C): 20, 30',
    'Rate (arbitrary units): 4, 8',
    '3. An investigation measures gas volume over time.',
    '(a) State the independent variable. [1]',
    '(b) Calculate the change in gas volume from 4 to 10 cubic centimetres. [2]',
    '4. Evaluate the method and suggest two improvements to its reliability. [8]',
]


def pdf_bytes(lines=LINES):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=700, height=842)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    escaped = [line.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)') for line in lines]
    stream.set_data(('BT /F1 10 Tf 40 790 Td 20 TL\n' + '\n'.join(f'({line}) Tj T*' for line in escaped) + '\nET').encode('latin-1'))
    page[NameObject('/Contents')] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes(lines=LINES, table=True):
    from docx import Document
    doc = Document()
    for index, line in enumerate(lines):
        if table and lines is LINES and index == 5:
            grid = doc.add_table(rows=3, cols=2)
            for row, values in zip(grid.rows, [('Temperature (C)', 'Rate'), ('20', '4'), ('30', '8')]):
                for cell, value in zip(row.cells, values):
                    cell.text = value
        if table and lines is LINES and index in {5, 6}:
            continue
        doc.add_paragraph(line)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


if __name__ == '__main__':
    target = Path(__file__).resolve().parents[1] / 'examples' / 'assessment_import'
    target.mkdir(parents=True, exist_ok=True)
    for extension, content in [('pdf', pdf_bytes()), ('docx', docx_bytes())]:
        path = target / ('fictional-test.' + extension)
        if not path.exists():
            path.write_bytes(content)
