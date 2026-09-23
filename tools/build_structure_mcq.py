"""Build a curated 100-question SL collection from the three classified banks.

All selected original diagrams are complete embedded PNGs, checked visually.
Adaptations isolate an in-scope prompt/context from a Paper 2/3 parent record;
their new choices and answer letters are explicitly distinguished from originals.
"""
from pathlib import Path
from collections import Counter
from copy import copy
from hashlib import sha256
from html import unescape
from io import BytesIO
import base64
import re
import struct
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
GUIDE = 'https://ibo.org/globalassets/new-structure/university-admission/pdfs/subject-guides/chemistry-guide.pdf'
SELECTION = {
    'S1.1': '166935 261576 502155 390669 799401 976176 877215',
    'S1.2': '130663 217048 166940 168292 328889 502157 390676 390851 801545 696383 976194 976210 874450',
    'S1.3': '130211 130212 217050 166941 328886 261581 499530 390852 386469 628208 799921 692297 696384 874455 877242',
    'S2.1': '261584 261847 502159 390687 386273 390861 799929 799930 692305 696393 874483 874484',
    'S2.2': '130216 130218 130607 130677 217061 217062 166945 166947 168298 168299 168302 328916 328919 328920 328922 261585 261587 217060 261850 499564 499566 390685 390865 390867 628212 799933 799934 692302 692308 696391 696392 696394 696954 976219 976221 976224 877255 877256 877266',
}
TITLES = {'S1.1': 'Particulate nature of matter', 'S1.2': 'The nuclear atom', 'S1.3': 'Electron configurations', 'S2.1': 'The ionic model', 'S2.2': 'The covalent model'}
# section, source paper, source URL ID, isolated source evidence, prompt,
# A/B/C/D choices, new correct letter, explanation
ADAPTED = [
    ('S1.1', 3, '634454', 'Recovering the evaporated solvent from extracted oil',
     'A student extracts oil from potato crisps using a volatile solvent. Which method allows the solvent to be collected separately from the oil?',
     ['Filtration alone', 'Evaporation into the room', 'Distillation using a condenser and receiver', 'Adding more crushed crisps'], 'C', 'The solvent vaporizes and is condensed into a separate receiver.'),
    ('S1.1', 3, '639914', 'Fractional distillation subprompt; excludes mole fraction and vapour-pressure calculations',
     'A mixture of miscible liquids is separated by fractional distillation. Which statement describes the separation?',
     ['Repeated vaporization and condensation separates components according to volatility.', 'The liquid with the highest boiling point is always collected first.', 'Covalent bonds in the liquids must be broken.', 'The liquids must react chemically with the glass beads.'], 'A', 'Repeated physical changes enrich the more volatile component in the vapour.'),
    ('S1.1', 2, '131432', 'Adapted from the stated melting point of magnesium chloride, 987 K',
     'Magnesium chloride melts at 987 K. What is this temperature in degrees Celsius, to the nearest whole degree?',
     ['260 °C', '987 °C', '1260 °C', '714 °C'], 'D', '987 − 273.15 = 713.85 °C, approximately 714 °C.'),
    ('S1.2', 2, '131432', 'Nuclear notation for magnesium-26',
     'Which pair gives the atomic number and mass number, respectively, of magnesium-26?',
     ['12 and 14', '12 and 26', '14 and 26', '26 and 12'], 'B', 'Magnesium has 12 protons; magnesium-26 has 26 nucleons.'),
    ('S1.2', 2, '167057', 'Numbers of subatomic particles in titanium-48',
     'A neutral titanium-48 atom has atomic number 22. How many protons, neutrons and electrons does it contain, respectively?',
     ['22, 26, 22', '22, 48, 22', '26, 22, 26', '22, 26, 20'], 'A', 'Protons = 22, neutrons = 48 − 22 = 26; a neutral atom has 22 electrons.'),
    ('S1.2', 2, '131432', 'Isotopic abundance calculation; source data restated as a table in text, not a mass spectrum',
     'A magnesium sample contains 78.60% magnesium-24, 10.11% magnesium-25 and 11.29% magnesium-26. What is its relative atomic mass to two decimal places?',
     ['24.00', '24.31', '25.00', '24.33'], 'D', '24 × 0.7860 + 25 × 0.1011 + 26 × 0.1129 = 24.3269.'),
    ('S1.3', 2, '167057', 'Electron configuration of titanium(II)',
     'Titanium has atomic number 22. What is the ground-state electron configuration of Ti²⁺?',
     ['[Ar]4s²', '[Ar]3d¹4s¹', '[Ar]3d²', '[Ar]3d²4s²'], 'C', 'The two 4s electrons are removed first from neutral titanium.'),
    ('S1.3', 2, '218103', 'Description of the hydrogen emission spectrum',
     'Which description of the hydrogen line emission spectrum is correct?',
     ['It contains every possible wavelength without gaps.', 'Its lines get closer together towards higher frequency.', 'Its lines get closer together towards longer wavelength.', 'Every line corresponds to an electron absorbing a photon.'], 'B', 'Discrete lines converge towards higher frequency and higher photon energy.'),
    ('S1.3', 2, '131432', 'New electron-configuration prompt based on the magnesium-26 context',
     'What is the full ground-state electron configuration of a neutral magnesium-26 atom (atomic number 12)?',
     ['1s²2s²2p⁶3s²', '1s²2s²2p⁶', '1s²2s²2p⁶3s¹', '1s²2s²2p⁶3s²3p²'], 'A', 'The neutral atom has 12 electrons; its neutron number does not change the configuration.'),
    ('S2.1', 2, '131432', 'Structure and bonding of magnesium oxide',
     'Which description of solid magnesium oxide is correct?',
     ['Separate MgO molecules held together by London forces', 'Magnesium and oxygen atoms sharing a sea of electrons', 'A giant lattice of Mg²⁺ and O²⁻ ions held by electrostatic attraction', 'A lattice of Mg⁺ and O⁻ ions held by hydrogen bonding'], 'C', 'MgO is a three-dimensional ionic lattice of oppositely charged ions.'),
    ('S2.1', 2, '387349', 'Structure and bonding in sodium oxide',
     'Which particles and attractions form the structure of solid sodium oxide, Na₂O?',
     ['Na atoms and O₂ molecules attracted by London forces', 'Na²⁺ and O⁻ ions in a 1:2 ratio', 'Na⁺ and O²⁻ ions in a 1:1 ratio', 'Na⁺ and O²⁻ ions in a 2:1 ratio, attracted electrostatically'], 'D', 'Two sodium ions balance the charge of each oxide ion in an ionic lattice.'),
    ('S2.1', 2, '167057', 'Type of bonding in potassium chloride',
     'Which interaction is responsible for the bonding in solid potassium chloride?',
     ['Attraction between oppositely charged potassium and chloride ions', 'Attraction between shared electron pairs and two chlorine nuclei', 'Hydrogen bonding between KCl molecules', 'Attraction between positive ions and delocalized electrons'], 'A', 'The source identifies KCl as ionic; its lattice is held by electrostatic attraction.'),
    ('S2.2', 3, '261927', 'Graphene structure subprompt; no band theory or mechanical calculations',
     'Which description explains why graphene is regarded as a two-dimensional material?',
     ['It consists of separate C₆ molecules.', 'It consists of a single sheet of covalently bonded carbon atoms.', 'It consists of a three-dimensional tetrahedral network.', 'It consists of an alternating lattice of positive and negative carbon ions.'], 'B', 'Graphene is one atom thick: a single covalent carbon sheet.'),
    ('S2.2', 3, '634454', 'Choice of a non-polar solvent to extract oil',
     'Why is a non-polar solvent suitable for extracting non-polar oil from potato crisps?',
     ['It ionizes all of the oil molecules.', 'It reacts with the oil to produce water.', 'Oil and solvent can interact through London dispersion forces.', 'It breaks every covalent bond in the oil.'], 'C', 'Non-polar solute and solvent can mix through compatible intermolecular interactions.'),
]

SUP = str.maketrans('0123456789+-−–=()ni', '⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁻⁻⁼⁽⁾ⁿⁱ')
SUB = str.maketrans('0123456789+-−–=()', '₀₁₂₃₄₅₆₇₈₉₊₋₋₋₌₍₎')


def math_text(node):
    kind = node.tag.split('}')[-1]
    children = [math_text(c) for c in node]
    if kind == 'msub':
        return children[0] + children[1].translate(SUB)
    if kind == 'msup':
        return children[0] + children[1].translate(SUP)
    if kind == 'msubsup':
        return children[0] + children[2].translate(SUP) + children[1].translate(SUB)
    if kind == 'mmultiscripts':
        split = next(i for i, c in enumerate(node) if c.tag.endswith('mprescripts'))
        before = children[split+1:]
        prefix = before[1].translate(SUP) + before[0].translate(SUB)
        after = ''.join(children[1:split]).translate(SUP)
        return prefix + children[0] + after
    if kind in {'mtr', 'mtd'}:
        return ''.join(children) + ('\n' if kind == 'mtr' else ' ')
    return (node.text or '').strip() + ''.join(children)


def readable(raw):
    raw = str(raw or '')
    image_count = 0
    def replace_image(match):
        nonlocal image_count
        image_count += 1
        return ' [See diagram / answer table in column E] '
    raw = re.sub(r'<img\b[^>]*>', replace_image, raw, flags=re.I)
    raw = re.sub(r'<math\b.*?</math>', lambda m: math_text(ET.fromstring(m[0])), raw, flags=re.S)
    raw = re.sub(r'<sup>(.*?)</sup>', lambda m: unescape(m[1]).translate(SUP), raw, flags=re.S)
    raw = re.sub(r'<sub>(.*?)</sub>', lambda m: unescape(m[1]).translate(SUB), raw, flags=re.S)
    raw = re.sub(r'</p>\s*,?\s*', '\n', raw)
    raw = re.sub(r'<br\s*/?>', '\n', raw)
    raw = unescape(re.sub(r'<[^>]+>', '', raw))
    raw = raw.strip().removeprefix('[').removesuffix(']').strip()
    return '\n'.join(re.sub(r'[^\S\n]+', ' ', line).strip(' ,') for line in raw.splitlines() if line.strip(' ,'))


class PreservedPNG(Image):
    """Embed original PNG bytes without re-encoding or a Pillow dependency."""
    def __init__(self, data):
        if data[:8] != b'\x89PNG\r\n\x1a\n' or b'IEND' not in data[-20:]:
            raise ValueError('Incomplete PNG')
        self.png = data
        self.width, self.height = struct.unpack('>II', data[16:24])
        self.format = 'png'

    def _data(self):
        return self.png


def load_sources():
    sources = {}
    for paper in (1, 2, 3):
        path = ROOT / f'output/exam_databases/Chem paper {paper} classified.xlsx'
        book = load_workbook(path, read_only=False, data_only=False)
        ws = book['Sheet1']
        headers = [c.value for c in ws[1]]
        records = {str(row[0].value): row for row in ws.iter_rows(min_row=2)}
        sources[paper] = (path, book, headers, records)
    return sources


def build(destination=None):
    sources = load_sources()
    hashes = {p: sha256(item[0].read_bytes()).hexdigest() for p, item in sources.items()}
    records = []
    for section, ids in SELECTION.items():
        for qid in ids.split():
            path, book, headers, rows = sources[1]
            source = dict(zip(headers, [c.value for c in rows[qid]]))
            raw = source['Question']
            images = re.findall(r'src=[\"\']data:image/png;base64,([^\"\']+)[\"\']', raw)
            if raw.count('<img') != len(images) or len(raw) >= 32767:
                raise ValueError(f'Question {qid}: incomplete or external image/content')
            text = readable(raw)
            if qid == '502157':
                text = 'What is represented by A in the nuclear notation ᴬZX²⁻ (A is the upper-left number and Z the lower-left number)?\nA. Number of electrons\nB. Number of neutrons\nC. Number of nucleons\nD. Number of protons'
            if qid == '696383':
                text = 'What is represented by “2−” in the nuclear notation for X²⁻?\nA. loss of electron\nB. gain of electron\nC. loss of proton\nD. gain of proton'
            if not images and not all(re.search(r'\b'+letter+r'\.', text) for letter in 'ABCD'):
                raise ValueError(f'Question {qid}: missing choices')
            answer = str(source['Markscheme']).strip()
            if answer not in 'ABCD' or len(answer) != 1:
                raise ValueError(f'Question {qid}: missing original key')
            records.append(dict(section=section, id=qid, paper=1, source_id=qid, text=text, answer=answer,
                                kind='Original MCQ', note='Formatting normalized; original wording and answer order retained.',
                                explanation='', images=[base64.b64decode(b) for b in images]))
        for j, (sec, paper, qid, evidence, prompt, options, answer, explanation) in enumerate(ADAPTED, 1):
            if sec == section:
                records.append(dict(section=sec, id=f'ADAPT_P{paper}_{qid}_{j:02}', paper=paper, source_id=qid,
                                    text=prompt+'\n'+'\n'.join(f'{a}. {v}' for a, v in zip('ABCD', options)),
                                    answer=answer, kind='Adapted to MCQ', note=evidence+'; new MCQ wording, distractors and key.',
                                    explanation=explanation, images=[]))
    assert len(records) == 100, len(records)
    assert len({r['id'] for r in records}) == 100
    assert len({r['text'] for r in records}) == 100
    out = Workbook()
    intro = out.active
    intro.title = 'Read me'
    intro.append(['100 MCQs — Structure 1.1, 1.2, 1.3, 2.1 and 2.2', ''])
    for row in [
        ('Scope', 'Current SL content (also suitable for HL core revision). Original source level is retained separately.'),
        ('Contents', '86 original Paper 1 MCQs; 14 adaptations from Paper 2/3. All questions carry one mark.'),
        ('Use', 'Questions contains prompts and embedded diagrams. AnswerKey is separate. Source sheets preserve the selected parent records.'),
        ('Syllabus review', 'Content checked against the official first-assessment-2025 guide; section-level mappings are assistant-reviewed, not teacher-confirmed.'),
        ('Guide', GUIDE),
        ('Excluded', 'HL resonance, expanded octets, hybridization, mass-spectrum interpretation and ionization-energy calculations; out-of-section content, duplicate exports, missing or truncated images.'),
        ('Adaptations', 'Each adaptation names the source record and isolated prompt/context. New distractors and keys are not original IB MCQs. No source-wide part mapping is claimed.'),
        ('Original formatting', 'HTML/MathML converted to readable notation. Complete source PNGs embedded unchanged. Raw source fields remain on Source_P1/P2/P3; source formulas are retained as literal text to avoid incorrect compacted-row references.'),
        ('Answers', 'Original Paper 1 key letters retained. Adapted answers have explanations. Teacher review is recommended before use as a scored assessment.'),
        ('Classification', 'QuestionTags preserves original tags only for original MCQs. Parent tags are not copied onto adaptations. MappingReview records proposed current sections.'),
        ('Coverage', 'Counts reflect available usable questions; this is a revision collection, not balanced coverage of every syllabus statement.'),
        ('Integration', 'Standalone question bank, not a complete Studentdash assessment workbook.'),
    ]:
        intro.append(row)
    for sec, count in Counter(r['section'] for r in records).items():
        intro.append((sec+' — '+TITLES[sec], count))
    for paper, digest in hashes.items():
        intro.append((f'Input Paper {paper} SHA256', digest))
    questions = out.create_sheet('Questions')
    questions.append(['Number', 'Structure', 'QuestionID', 'Question and options', 'Diagram / answer table', 'Current level', 'Type', 'Source paper', 'Source ID', 'Source reference', 'Source level', 'Marks'])
    answers = out.create_sheet('AnswerKey')
    answers.append(['Number', 'QuestionID', 'Structure', 'Answer', 'Explanation for adapted question', 'Type'])
    review = out.create_sheet('MappingReview')
    review.append(['QuestionID', 'ProposedSection', 'MappingSource', 'ReviewStatus', 'Selection / adaptation note'])
    tags = out.create_sheet('QuestionTags')
    tags.append(['QuestionID', 'Category', 'Tag', 'Source', 'Confidence'])
    source_ids = {p: set() for p in sources}
    for n, r in enumerate(records, 1):
        p, book, headers, rows = sources[r['paper']]
        src = dict(zip(headers, [c.value for c in rows[r['source_id']]]))
        source_ids[r['paper']].add(r['source_id'])
        questions.append([n, r['section'], r['id'], r['text'], '', 'SL / HL core', r['kind'], r['paper'], r['source_id'], src['Reference code'], src['Level'], 1])
        row = n+1
        questions.row_dimensions[row].height = min(409, max(115, (len(r['text'])//95 + r['text'].count('\n')+2)*15))
        for png in r['images']:
            image = PreservedPNG(png)
            scale = min(1, 420/image.width, 470/image.height)
            image.width *= scale
            image.height *= scale
            questions.add_image(image, f'E{row}')
            questions.row_dimensions[row].height = max(questions.row_dimensions[row].height, image.height*.75+12)
        answers.append([n, r['id'], r['section'], r['answer'], r['explanation'], r['kind']])
        review.append([r['id'], r['section'], 'assistant content review', 'proposed — teacher confirmation pending', r['note']])
    for row in list(sources[1][1]['QuestionTags'].values)[1:]:
        if str(row[0]) in source_ids[1]:
            tags.append(row)
    for paper, (_, book, headers, rows) in sources.items():
        ws = out.create_sheet(f'Source_P{paper}')
        ws.append(headers)
        for qid in sorted(source_ids[paper]):
            cells = rows[qid]
            for col, cell in enumerate(cells, 1):
                dest = ws.cell(ws.max_row+1 if col == 1 else ws.max_row, col, cell.value)
                # Style IDs belong to the source workbook; use component copies.
                dest.font, dest.fill, dest.border = copy(cell.font), copy(cell.fill), copy(cell.border)
                dest.alignment, dest.protection = copy(cell.alignment), copy(cell.protection)
                dest.number_format = cell.number_format
        # Original formula strings kept verbatim as text: source row references
        # do not refer to this compacted sheet. No formula is executed/rebased.
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'
        ws.sheet_state = 'hidden'
    for ws in out:
        if ws.title.startswith('Source_'):
            continue
        ws.freeze_panes = 'D2' if ws == questions else 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='183D52')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name='Calibri', size=11)
                c.alignment = Alignment(wrap_text=True, vertical='top')
                if c.row % 2 == 0:
                    c.fill = PatternFill('solid', fgColor='EEF4F6')
        for col in range(1, ws.max_column+1):
            ws.column_dimensions[get_column_letter(col)].width = 22
    intro.column_dimensions['A'].width = 36
    intro.column_dimensions['B'].width = 110
    for row in range(2, intro.max_row+1):
        intro.row_dimensions[row].height = 42
    for col, width in {'A':9, 'B':13, 'C':29, 'D':96, 'E':62, 'F':17, 'G':22, 'H':14, 'I':14, 'J':25, 'K':14, 'L':10}.items():
        questions.column_dimensions[col].width = width
    answers.column_dimensions['E'].width = 95
    review.column_dimensions['D'].width = 42
    review.column_dimensions['E'].width = 110
    intro['B6'].hyperlink = GUIDE
    intro['B6'].style = 'Hyperlink'
    out.active = 1
    destination = Path(destination) if destination else ROOT/'output/exam_databases/Structure 1.1-1.3 and 2.1-2.2 - 100 MCQ.xlsx'
    if destination.exists():
        raise FileExistsError(f'Refusing to overwrite {destination}')
    out.save(destination)
    for p, item in sources.items():
        item[1].close()
        assert sha256(item[0].read_bytes()).hexdigest() == hashes[p]
    check = load_workbook(destination, read_only=True)
    assert check['Questions'].max_row == 101
    assert check['AnswerKey'].max_row == 101
    check.close()
    print(destination)
    print(dict(Counter(r['section'] for r in records)))
    print(dict(Counter(r['paper'] for r in records)))
    print('Originals', sum(r['kind']=='Original MCQ' for r in records), 'Embedded images', sum(len(r['images']) for r in records))


if __name__ == '__main__':
    build()
