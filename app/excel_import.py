"""Read-only adapter for Overview + copied, date-named Assessment blueprint sheets."""
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import math
import re
import unicodedata
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook

from .storage import ValidationError, course_id, now


def identity(value):
    return ' '.join(unicodedata.normalize('NFKC', str(value or '')).casefold().split())


def number(value, label, positive=False):
    if isinstance(value, bool) or not isinstance(value, (float, int)) or not math.isfinite(value) or value < 0 or (positive and value == 0):
        raise ValidationError(f'{label}: expected a finite {"positive" if positive else "nonnegative"} number.')
    return float(value)


def open_excel(content, data_only=False):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024 or len(archive.infolist()) > 5000:
                raise ValidationError('Workbook expands beyond the 80 MB safety limit.')
        book = load_workbook(BytesIO(content), data_only=data_only)
        if any(s.max_row > 20000 or s.max_column > 1000 or s.max_row * s.max_column > 2000000 for s in book):
            book.close()
            raise ValidationError('Workbook dimensions are too large. Remove unused formatted rows/columns in a copy.')
        return book
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        if isinstance(exc, ValidationError):
            raise
        raise ValidationError('Cannot read this file as an Excel .xlsx workbook.') from exc


def score(value, maximum, label):
    if value is None or str(value).strip() == '':
        return dict(status='pending', score=None)
    states = {'a': 'absent', 'x': 'absent', 'frav': 'absent', 'absent': 'absent',
              'e': 'exempt', 'exempt': 'exempt', 'm': 'missing', 'missing': 'missing',
              'p': 'pending', 'pending': 'pending'}
    state = states.get(str(value).strip().casefold())
    if state:
        return dict(status=state, score=None)
    mark = number(value, label)
    if mark > maximum:
        raise ValidationError(f'{label}: score exceeds maximum {maximum:g}.')
    return dict(status='graded', score=mark)


def read_grading_workbook(content):
    formulas = open_excel(content)
    values = open_excel(content, data_only=True)
    warnings = []
    try:
        if 'Overview' not in formulas:
            raise ValidationError('Expected an Overview sheet with Name in A1 and students from A2.')

        def cell(sheet, row, col, trail=()):
            raw = formulas[sheet].cell(row, col)
            cached = values[sheet].cell(row, col)
            if raw.data_type == 'e' or cached.data_type == 'e':
                raise ValidationError(f'{sheet}!{raw.coordinate}: Excel error. Correct it and save in Excel.')
            if raw.data_type != 'f':
                return raw.value
            # A direct cell link (the supplied blueprint's name column) is safe to resolve.
            match = re.fullmatch(r"=(?:'([^']+)'|([A-Za-z0-9_ ]+))!\$?([A-Z]+)\$?(\d+)", raw.value)
            if match and len(trail) < 10:
                from openpyxl.utils.cell import column_index_from_string
                target = (match[1] or match[2], int(match[4]), column_index_from_string(match[3]))
                if target[0] in formulas.sheetnames and target not in trail:
                    return cell(*target, trail=trail + (target,))
            if cached.value is not None:
                return cached.value
            raise ValidationError(f'{sheet}!{raw.coordinate}: formula has no cached result. Recalculate and save in Excel, then upload again.')

        if identity(cell('Overview', 1, 1)) not in {'name', 'student', 'student name'}:
            raise ValidationError('Overview!A1 must identify the student-name column (Name).')
        headers = {identity(c.value): c.column for c in formulas['Overview'][1] if c.value is not None}
        roster, seen = [], set()
        for r in range(2, formulas['Overview'].max_row + 1):
            name = cell('Overview', r, 1)
            if name is None or not str(name).strip():
                continue
            name = str(name).strip()
            external = cell('Overview', r, headers['studentid']) if 'studentid' in headers else None
            key = 'id:' + str(external).strip() if external is not None else 'name:' + identity(name)
            if key in seen or any(identity(s['name']) == identity(name) for s in roster):
                raise ValidationError(f'Overview!A{r}: duplicate/ambiguous student. Use distinct names and stable StudentID values.')
            seen.add(key)
            roster.append(dict(key=key, name=name, level=str(cell('Overview', r, headers['level']) or '') if 'level' in headers else ''))
        if not roster:
            raise ValidationError('Overview contains no students.')
        known = {identity(s['name']): s['key'] for s in roster}
        boundaries = []
        if 'Grade boundaries' in formulas:
            for r in range(2, formulas['Grade boundaries'].max_row + 1):
                grade, percent = cell('Grade boundaries', r, 1), cell('Grade boundaries', r, 2)
                if grade is None or isinstance(grade, str):
                    continue  # FRAV is an absence label, not a grade.
                if isinstance(grade, bool) or int(grade) != grade:
                    raise ValidationError('Grade boundaries: grades must be integers.')
                threshold = number(percent, f'Grade boundaries!B{r}')
                if threshold > 100:
                    raise ValidationError('Grade boundaries percentages must be between 0 and 100.')
                boundaries.append(dict(grade=int(grade), percent=threshold))
            boundaries.sort(key=lambda b: b['percent'])
            if boundaries and (boundaries[0]['percent'] != 0 or len({b['grade'] for b in boundaries}) != len(boundaries) or
                               any(a['percent'] >= b['percent'] or a['grade'] >= b['grade'] for a, b in zip(boundaries, boundaries[1:]))):
                raise ValidationError('Grade boundaries must be unique and strictly increasing, starting at 0%.')
        if not boundaries:
            warnings.append('No grade boundaries: percentages will be shown without suggested grades.')
        assessments = {}
        ignored = []
        for sheet in formulas:
            if sheet.title in {'Overview', 'Grade boundaries', 'Assessment blueprint'}:
                continue
            day = None
            for fmt in ('%d.%m.%Y', '%Y-%m-%d'):
                try:
                    day = datetime.strptime(sheet.title, fmt).date().isoformat()
                    break
                except ValueError:
                    pass
            if day is None:
                ignored.append(sheet.title)
                continue
            if day in assessments:
                raise ValidationError(f'Two sheets identify the same assessment date {day}.')
            # Administrative columns begin at Marks; numeric-looking columns after it are never questions.
            total_cols = [c.column for c in sheet[1] if identity(c.value) in {'marks', 'total', 'total marks'}]
            if len(total_cols) != 1:
                raise ValidationError(f'{sheet.title}: expected exactly one Marks total column in row 1.')
            stop = total_cols[0]
            questions = {}
            for c in range(2, stop):
                label, maximum = cell(sheet.title, 1, c), cell(sheet.title, 2, c)
                if label is None and maximum is None:
                    continue
                label = str(label or '').strip()
                if not label or len(label) > 80 or label.casefold() in {q.casefold() for q in questions}:
                    raise ValidationError(f'{sheet.title}: empty or duplicate question label in column {c}.')
                questions[label] = dict(id=label, marks=number(maximum, f'{sheet.title} row 2 column {c}', True), column=c)
            if not questions:
                raise ValidationError(f'{sheet.title}: no question columns before Marks.')
            absence = next((c.column for c in sheet[1] if identity(c.value) in {'x for absence', 'absence', 'absent', 'fravær'}), None)
            results = {}
            for r in range(3, sheet.max_row + 1):
                name = cell(sheet.title, r, 1)
                if name is None or not str(name).strip() or identity(name) in {'sum', 'task', 'total', 'name'}:
                    continue
                key = known.get(identity(name))
                if key is None:
                    raise ValidationError(f'{sheet.title}!A{r}: student not found in Overview. Check the name/link.')
                if key in results:
                    raise ValidationError(f'{sheet.title}!A{r}: duplicate student row.')
                flag = cell(sheet.title, r, absence) if absence else None
                absent = identity(flag) in {'x', 'a', 'absent', 'frav', 'yes', '1'}
                if flag is not None and str(flag).strip() and not absent:
                    raise ValidationError(f'{sheet.title} row {r}: unrecognized absence marker.')
                marks = {qid: score(cell(sheet.title, r, q['column']), q['marks'], f'{sheet.title} row {r}, {qid}') for qid, q in questions.items()}
                if absent:
                    if any(m['status'] == 'graded' for m in marks.values()):
                        warnings.append(f'{sheet.title} row {r}: absence marker overrides entered question marks.')
                    marks = {qid: dict(status='absent', score=None) for qid in questions}
                results[key] = marks
                # Cached totals are only a consistency check when evidence is complete.
                if all(m['status'] in {'graded', 'exempt'} for m in marks.values()):
                    total_cell = values[sheet.title].cell(r, stop)
                    total = total_cell.value
                    if isinstance(total, (int, float)) and not isinstance(total, bool) and math.isfinite(total):
                        earned = sum(m['score'] for m in marks.values() if m['status'] == 'graded')
                        if abs(earned - total) > 0.00001:
                            warnings.append(f'{sheet.title} row {r}: workbook total differs from question marks; inspect formulas. Question marks are retained.')
            for s in roster:
                if s['key'] not in results:
                    warnings.append(f'{sheet.title}: an Overview student has no score row; results remain pending.')
                    results[s['key']] = {qid: dict(status='pending', score=None) for qid in questions}
            assessments[day] = dict(id=day, date=day, name=sheet.title, questions={qid: {'id': qid, 'marks': q['marks']} for qid, q in questions.items()}, results=results)
        if not assessments:
            warnings.append('No dated assessment sheets found. Assessment blueprint is intentionally excluded; copy it and name the copy DD.MM.YYYY when entering a test.')
        if ignored:
            warnings.append('Ignored non-assessment sheets: ' + ', '.join(ignored))
        warnings.append('Formulas are never calculated or written. Cached score formulas must have been recalculated and saved in Excel; cached freshness cannot be verified. Totals and grades are recomputed from complete question evidence; missing marks stay pending.')
        return dict(students=roster, assessments=assessments, boundaries=boundaries, warnings=warnings,
                    fingerprint=sha256(content).hexdigest())
    finally:
        formulas.close()
        values.close()


def structural_changes(doc, incoming):
    changes = []
    if not doc:
        return changes
    for aid, old in doc['assessments'].items():
        new = incoming['assessments'].get(aid)
        if new is None:
            if old.get('active', True):
                changes.append(f'{aid}: absent from upload; retain history/classifications but exclude from current dashboards.')
            continue
        for qid, q in old['questions'].items():
            if not q.get('active', True):
                continue
            if qid not in new['questions']:
                changes.append(f'{aid} / {qid}: removed. Existing classification will be retained as history, not applied elsewhere.')
            elif q['marks'] != new['questions'][qid]['marks']:
                changes.append(f'{aid} / {qid}: maximum changed from {q["marks"]:g} to {new["questions"][qid]["marks"]:g}. Classification needs review and will be retired.')
        for qid in new['questions']:
            if qid in old['questions'] and not old['questions'][qid].get('active', True):
                changes.append(f'{aid} / {qid}: retired identifier reused; old classifications will not be reattached.')
    return changes


def import_results(store, key, incoming, revision, profile, curriculum, confirm=False, mappings=None, default_level=''):
    course_id(key)
    mappings = mappings or {}

    def update(doc, db):
        if structural_changes(doc, incoming) and not confirm:
            raise ValidationError('Question structure changed. Review and explicitly confirm the changes first.')
        if doc is None:
            doc = dict(id=key, profile=deepcopy(profile), curriculum=deepcopy(curriculum), students={}, assessments={}, boundaries=[], default_level=default_level)
        by_source = {s['source_key']: sid for sid, s in doc['students'].items()}
        resolved = {}
        for student in incoming['students']:
            source = student['key']
            sid = mappings.get(source) or by_source.get(source)
            if sid and sid not in doc['students']:
                raise ValidationError('Unknown mapped StudentID.')
            if sid in resolved.values():
                raise ValidationError('Two imported students map to the same StudentID.')
            if not sid:
                sid = str(db.execute('SELECT next_student FROM sequence WHERE id=1').fetchone()[0])
                db.execute('UPDATE sequence SET next_student=next_student+1 WHERE id=1')
            resolved[source] = sid
            doc['students'][sid] = dict(student, id=sid, source_key=source, active=True)
        for sid, student in doc['students'].items():
            student['active'] = sid in resolved.values()
        doc['boundaries'] = incoming['boundaries']
        doc['source_fingerprint'] = incoming['fingerprint']
        for aid, item in doc['assessments'].items():
            item['active'] = aid in incoming['assessments']
        for aid, item in incoming['assessments'].items():
            old = doc['assessments'].get(aid, dict(questions={}))
            for qid, question in old['questions'].items():
                new = item['questions'].get(qid)
                if new is None or new['marks'] != question['marks'] or not question.get('active', True):
                    if question.get('classification'):
                        question.setdefault('retired', []).append(dict(classification=question.pop('classification'), marks=question['marks'], retired=now()))
                question['active'] = new is not None
            for qid, q in item['questions'].items():
                old['questions'].setdefault(qid, {})
                old['questions'][qid].update(q, active=True)
            old.update(id=aid, date=item['date'], name=item['name'], active=True,
                       results={resolved[source]: marks for source, marks in item['results'].items()})
            doc['assessments'][aid] = old
        return doc

    return store.commit(key, revision, update, 'results', dict(fingerprint=incoming['fingerprint'],
                        assessments=list(incoming['assessments']), warnings=incoming['warnings'],
                        mappings=mappings, structure_confirmed=confirm))
