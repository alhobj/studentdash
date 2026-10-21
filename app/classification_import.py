"""Explicit, previewed classification changes; no inference or silent replacement."""
import csv
from copy import deepcopy
from hashlib import sha256
from io import BytesIO, StringIO
import json

from openpyxl import Workbook

from .excel_import import open_excel
from .storage import ValidationError, now

IDENTITY = ['ClassID', 'AssessmentID', 'QuestionID']
OPTIONAL = ['Level', 'SyllabusCode', 'Topic', 'Subtopic']
CLEAR = '[clear]'


def node_index(curriculum):
    nodes = curriculum.get('nodes', [])
    index = {}
    codes = set()
    for node in nodes:
        if not isinstance(node, dict) or not isinstance(node.get('id'), str) or not node['id'] or node['id'] in index or not isinstance(node.get('label'), str) or not node['label'].strip():
            raise ValidationError('SyllabusNodes needs unique NodeID values and a Label.')
        code = node.get('code')
        if code and code in codes:
            raise ValidationError('Syllabus codes must be unique.')
        if code:
            codes.add(code)
        index[node['id']] = node
    for key in index:
        seen, cursor = set(), key
        while cursor:
            if cursor in seen or cursor not in index:
                raise ValidationError('SyllabusNodes has an unknown parent or a cycle.')
            seen.add(cursor)
            cursor = index[cursor].get('parent')
    return index


def headers(profile):
    return IDENTITY + OPTIONAL + list(profile['categories'])


def text(value):
    return '' if value is None else str(value).strip()


def tags(value):
    return list(dict.fromkeys(s.strip() for s in text(value).split(';') if s.strip()))


def rows_from_table(rows):
    rows = iter(rows)
    heading = [text(v) for v in next(rows, [])]
    if len(set(heading)) != len(heading) or not set(IDENTITY) <= set(heading):
        raise ValidationError('Classification headers must be unique and include ClassID, AssessmentID and QuestionID.')
    result = []
    for index, values in enumerate(rows, 2):
        if not any(v is not None and text(v) for v in values):
            continue
        if len(values) > len(heading) and any(text(v) for v in values[len(heading):]):
            raise ValidationError(f'Classification row {index}: too many cells.')
        result.append(dict(zip(heading, list(values) + [''] * max(0, len(heading) - len(values)))))
    if not result:
        raise ValidationError('The classification file contains no question rows.')
    if len(result) > 10000:
        raise ValidationError('At most 10000 classification rows may be imported at once.')
    return result, heading


def parse_file(content, filename):
    nodes = []
    if filename.lower().endswith('.csv'):
        try:
            rows, heading = rows_from_table(csv.reader(StringIO(content.decode('utf-8-sig'))))
        except (UnicodeError, csv.Error) as exc:
            raise ValidationError('CSV must be UTF-8, comma-delimited. Use semicolons between multiple tags in one cell.') from exc
    elif filename.lower().endswith('.xlsx'):
        book = open_excel(content)
        try:
            if 'Classifications' not in book:
                raise ValidationError('The workbook needs a Classifications sheet. Download the template.')
            for sheet in book:
                if any(c.data_type in {'f', 'e'} for row in sheet for c in row):
                    raise ValidationError('Classification files must contain values, not formulas or Excel errors.')
            rows, heading = rows_from_table(book['Classifications'].values)
            if 'SyllabusNodes' in book:
                table = iter(book['SyllabusNodes'].values)
                node_headers = [text(v) for v in next(table, [])]
                if node_headers != ['NodeID', 'ParentID', 'Label', 'Code']:
                    raise ValidationError('SyllabusNodes headers must be NodeID, ParentID, Label, Code.')
                for row in table:
                    if any(text(v) for v in row):
                        item = dict(zip(node_headers, map(text, row)))
                        nodes.append(dict(id=item.get('NodeID', ''), parent=item.get('ParentID') or None,
                                          label=item.get('Label', ''), code=item.get('Code', '')))
        finally:
            book.close()
    else:
        raise ValidationError('Choose a classification .xlsx or UTF-8 .csv file.')
    return dict(rows=rows, headers=heading, nodes=nodes, fingerprint=sha256(content).hexdigest())


def preview(doc, parsed):
    allowed = set(headers(doc['profile']))
    unknown = set(parsed['headers']) - allowed
    if unknown:
        raise ValidationError('Unknown classification columns: ' + ', '.join(sorted(unknown)))
    curriculum = deepcopy(doc['curriculum'])
    existing = node_index(curriculum)
    added, seen_nodes = [], set()
    for node in parsed['nodes']:
        if node['id'] in seen_nodes:
            raise ValidationError('Duplicate SyllabusNodes row.')
        seen_nodes.add(node['id'])
        old = existing.get(node['id'])
        if old:
            if any(text(old.get(k)) != text(node.get(k)) for k in ('parent', 'label', 'code')):
                raise ValidationError('Existing syllabus nodes cannot be redefined by an import. Add a new node ID/version instead.')
        else:
            curriculum.setdefault('nodes', []).append(node)
            added.append(node)
    index = node_index(curriculum)
    code_map = {node.get('code'): key for key, node in index.items() if node.get('code')}
    changes, conflicts, seen = [], [], set()
    for rownum, row in enumerate(parsed['rows'], 2):
        row = {k: text(v) for k, v in row.items()}
        key, aid, qid = (row.get(c, '') for c in IDENTITY)
        label = f'Row {rownum}: {aid} / {qid}'
        if key != doc['id']:
            raise ValidationError(f'{label}: unknown ClassID {key}. Select the matching class.')
        if aid not in doc['assessments'] or not doc['assessments'][aid].get('active', True):
            raise ValidationError(f'{label}: unknown or inactive AssessmentID.')
        question = doc['assessments'][aid]['questions'].get(qid)
        if not question or not question.get('active', True):
            raise ValidationError(f'{label}: unknown or retired QuestionID. Include any * prefix exactly.')
        if (aid, qid) in seen:
            raise ValidationError(f'{label}: duplicate classification row; combine tags with semicolons.')
        seen.add((aid, qid))
        old = deepcopy(question.get('classification', {}))
        new = deepcopy(old)
        new.setdefault('tags', {})
        for category, vocabulary in doc['profile']['categories'].items():
            value = row.get(category, '')
            if not value:
                continue
            proposed = [] if value == CLEAR else tags(value)
            if set(proposed) - set(vocabulary):
                raise ValidationError(f'{label}: unknown {category} tag(s): ' + ', '.join(sorted(set(proposed) - set(vocabulary))))
            new['tags'][category] = proposed
        for column, field in [('Level', 'level'), ('Topic', 'topic'), ('Subtopic', 'subtopic')]:
            value = row.get(column, '')
            if value:
                if column == 'Level' and value != CLEAR and value not in doc['profile'].get('levels', []):
                    raise ValidationError(f'{label}: Level must be one of the course profile levels.')
                new[field] = '' if value == CLEAR else value
        value = row.get('SyllabusCode', '')
        if value:
            selected = []
            for code in ([] if value == CLEAR else tags(value)):
                nid = code if code in index else code_map.get(code)
                if nid is None:
                    raise ValidationError(f'{label}: unknown syllabus code {code}. Add an explicit node and parent in SyllabusNodes; no hierarchy is guessed.')
                if nid not in selected:
                    selected.append(nid)
            new['nodes'] = selected
        for column in ('Topic', 'Subtopic'):
            if len(row.get(column, '')) > 300:
                raise ValidationError(f'{label}: {column} is too long.')
        # Metadata does not affect semantic idempotence.
        old_clean = {k: v for k, v in old.items() if k not in {'reviewed', 'source'}}
        new_clean = {k: v for k, v in new.items() if k not in {'reviewed', 'source'}}
        if not old and new_clean == {'tags': {}}:
            continue
        if old_clean != new_clean:
            replaced = []
            for category, values in new.get('tags', {}).items():
                previous = old.get('tags', {}).get(category, [])
                if previous and previous != values:
                    replaced.append(category)
            for field in ('level', 'topic', 'subtopic', 'nodes'):
                if old.get(field) and old.get(field) != new.get(field):
                    replaced.append(field)
            if replaced:
                conflicts.append(f'{aid} / {qid}: replace ' + ', '.join(replaced))
            changes.append(dict(assessment=aid, question=qid, before=old, after=new))
    return dict(changes=changes, conflicts=conflicts, added_nodes=added, curriculum=curriculum,
                fingerprint=parsed['fingerprint'], rows=len(parsed['rows']))


def import_classifications(store, key, parsed, revision, replace=False):
    def update(doc, db):
        proposal = preview(doc, parsed)
        if proposal['conflicts'] and not replace:
            raise ValidationError('Existing classifications are protected. Review and explicitly confirm replacements.')
        doc['curriculum'] = proposal['curriculum']
        for change in proposal['changes']:
            q = doc['assessments'][change['assessment']]['questions'][change['question']]
            q['classification'] = dict(change['after'], source='teacher-confirmed import', reviewed=now())
        return doc
    proposal = preview(store.read(key), parsed)
    return store.commit(key, revision, update, 'classifications',
                        dict(fingerprint=parsed['fingerprint'], changes=proposal['changes'],
                             added_nodes=proposal['added_nodes'], replacements_confirmed=replace))


def template(doc, scaffold=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Classifications'
    heading = headers(doc['profile'])
    sheet.append(heading)
    if scaffold:
        for aid, a in sorted(doc['assessments'].items()):
            if not a.get('active', True):
                continue
            for qid, q in a['questions'].items():
                if q.get('active', True):
                    c = q.get('classification', {})
                    values = dict(ClassID=doc['id'], AssessmentID=aid, QuestionID=qid,
                                  Level=c.get('level', ''), SyllabusCode='; '.join(c.get('nodes', [])),
                                  Topic=c.get('topic', ''), Subtopic=c.get('subtopic', ''))
                    values.update({k: '; '.join(v) for k, v in c.get('tags', {}).items()})
                    sheet.append([values.get(h, '') for h in heading])
    vocabulary = book.create_sheet('Vocabulary')
    vocabulary.append(['Category', 'Allowed value'])
    for category, values in doc['profile']['categories'].items():
        for value in values:
            vocabulary.append([category, value])
    for level in doc['profile'].get('levels', []):
        vocabulary.append(['Level', level])
    nodes = book.create_sheet('SyllabusNodes')
    nodes.append(['NodeID', 'ParentID', 'Label', 'Code'])
    for node in doc['curriculum'].get('nodes', []):
        nodes.append([node['id'], node.get('parent'), node['label'], node.get('code')])
    instructions = book.create_sheet('Instructions')
    for line in [
        'One row per ClassID + AssessmentID + QuestionID. Preserve * prefixes.',
        'AssessmentID is the ISO date YYYY-MM-DD. Use semicolons between tags in a cell.',
        'Blank cells preserve existing values; unknown new classifications remain unknown.',
        'Use [clear] to explicitly clear a field. Replacements require teacher confirmation.',
        'Use only Vocabulary values. Vocabulary is explanatory; edit profiles/*.json for new classes.',
        'Use SyllabusNodes NodeID or Code. Add specific statements with explicit ParentID; do not redefine existing nodes.',
        'Never include student names, scores, answer keys or private workbook content in a ChatGPT prompt.',
    ]:
        instructions.append([line])
    # Avoid spreadsheet formula injection from imported labels or teacher annotations.
    for tab in book:
        tab.freeze_panes = 'A2'
        for row in tab:
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = 's'
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def prompt(doc):
    return ('Classify the attached assessment questions, not student results. Complete the Classifications sheet '
            'of this Studentdash template, preserving ClassID, AssessmentID and QuestionID exactly. '
            'Keep * prefixes. Use only the Vocabulary values and separate multiple tags with semicolons. '
            'Leave uncertain fields blank; do not infer unseen question content. Use existing SyllabusNodes '
            'IDs/codes. If a specific statement is absent, add its explicit node, label, code and parent '
            'without changing existing nodes. Return an .xlsx file with the original sheet headers. '
            'Do not include learner names, scores or personal information. ClassID: ' + doc['id'])
